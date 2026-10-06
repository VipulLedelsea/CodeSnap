"""Edge cases of the line-by-line analysis: awkward quotes and line numbers, broken model replies, failures during a
run, invalidation, one job at a time, every rescan trigger, and a report that never mentions how the code was read."""
import re
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest

from conftest import need_fixture

from core import deepdive as DD
from core.model import ProgramStore

from test_deepdive import FakeClient, fact, tool_msg

RUN1 = Path(__file__).resolve().parent / "fixtures" / "run1_program"
CODE = "\n".join(["PROGRAM-ID. X.",                      # 1
                  "    MOVE A TO B",                     # 2
                  "    IF WS-FLAG = 'Y'",                # 3
                  "       PERFORM PAY-IT",               # 4
                  "    END-IF",                          # 5
                  "    MOVE A TO B",                     # 6
                  "    COMPUTE WS-NET = WS-GROSS -",     # 7
                  "            WS-HOLDBACK",             # 8
                  "    STOP RUN."])                      # 9


# ── quotes and line numbers ────────────────────────────────────────────────────────────────────────────────────

def test_quote_spanning_two_lines_is_accepted_with_both_lines():
    kept, *_ = DD.check_facts([fact("Net is gross minus holdback", [7, 8], "WS-GROSS - WS-HOLDBACK")], CODE)
    assert kept and kept[0]["lines"] == [7, 8]


def test_multi_line_quote_with_invented_tail_is_rejected():
    _, _, rej, _ = DD.check_facts([fact("x", [7, 8], "WS-GROSS - WS-HOLDBACK - WS-FEE")], CODE)
    assert rej and rej[0]["why"] == "quote not found in the file"


def test_repeated_quote_picks_the_occurrence_nearest_the_cited_lines():
    kept, corrected, *_ = DD.check_facts([fact("second move", [9, 9], "MOVE A TO B"), fact("first move", [1, 1], "MOVE A TO B")], CODE)
    assert [f["lines"] for f in kept] == [[6, 6], [2, 2]] and corrected == 2


def test_quote_on_the_cited_line_is_not_moved():
    kept, corrected, *_ = DD.check_facts([fact("second move", [6, 6], "MOVE A TO B")], CODE)
    assert kept[0]["lines"] == [6, 6] and corrected == 0


@pytest.mark.parametrize("lines", [[99, 120], [8, 7], ["7", "8"], [0], []])
def test_out_of_range_reversed_text_or_missing_line_numbers_are_handled(lines):
    kept, _, rejected, _ = DD.check_facts([fact("Net is gross minus holdback", lines, "WS-GROSS -")], CODE)
    assert not rejected and kept[0]["lines"][0] == 7


def test_malformed_facts_are_rejected_not_crashing():
    _, _, rej, _ = DD.check_facts([None, "text", {"statement": "no category", "quote": "MOVE A", "lines": [2]},
                                   fact("bad category", [2, 2], "MOVE A", cat="gossip"), fact("short quote", [2, 2], "A")], CODE)
    assert len(rej) == 5


def test_partly_unreadable_span_is_kept_fully_unreadable_is_held_back():
    kept, _, _, unver = DD.check_facts([fact("n", [7, 8], "WS-GROSS - WS-HOLDBACK"), fact("m", [8, 8], "WS-HOLDBACK")],
                                       CODE, bad=[8])
    assert [f["statement"] for f in kept] == ["n"] and [f["statement"] for f in unver] == ["m"]


# ── broken model replies and failures ──────────────────────────────────────────────────────────────────────────

@pytest.fixture
def store(tmp_path):
    need_fixture(RUN1)
    shutil.copytree(RUN1, tmp_path / "live-test-09-29")
    st = ProgramStore.open("live-test-09-29", root=tmp_path)
    yield st
    st.close()


def art(store, name):
    return next(a for a in store.artifacts() if a["name"] == name)


INQ = {"purpose": "Inquiry.", "facts": [fact("Receives map AIDMAP1", [18, 18], "RECEIVE MAP('AIDMAP1')", cat="ui"),
                                        fact("Links to AIDAUDIT", [25, 25], "LINK PROGRAM('AIDAUDIT')", cat="interface")]}


@pytest.mark.parametrize("review", [{"verdicts": [{"id": 1, "verdict": "supported"}]},        # id 2 missing
                                    {"verdicts": [{"id": "x", "verdict": "??"}, "junk"]},       # malformed
                                    {}])                                                       # empty
def test_incomplete_or_malformed_review_keeps_facts_as_not_reviewed(store, review):
    res = DD.analyse_file(store, FakeClient(INQ, review), art(store, "AIDINQ.cbl"))
    assert len(res["facts"]) == 2
    assert all(f["review"] in ("supported", "not reviewed") for f in res["facts"])


def test_text_only_json_reply_is_accepted():
    msg = SimpleNamespace(content=[SimpleNamespace(type="text", text='Here it is: {"purpose": "p", "facts": []}')])
    assert DD._tool(msg, "record_analysis") == {"purpose": "p", "facts": []}
    assert DD._tool(SimpleNamespace(content=[SimpleNamespace(type="text", text="sorry")]), "record_analysis") is None


class NoThinking(FakeClient):
    def create(self, **kw):
        if "thinking" in kw:
            self.calls.append(kw)
            raise TypeError("create() got an unexpected keyword argument 'thinking'")
        return super().create(**kw)


def test_falls_back_when_extended_reasoning_is_unavailable(store):
    client = NoThinking(INQ, {"verdicts": [{"id": 1, "verdict": "supported"}, {"id": 2, "verdict": "supported"}]})
    res = DD.analyse_file(store, client, art(store, "AIDINQ.cbl"))
    assert len(res["facts"]) == 2
    assert client.calls[1]["tool_choice"] == {"type": "tool", "name": "record_analysis"}


class Unrelated(FakeClient):
    def create(self, **kw):
        raise RuntimeError("connection reset by peer")


def test_unrelated_errors_are_not_swallowed_by_the_fallback(store):
    with pytest.raises(RuntimeError):
        DD.analyse_file(store, Unrelated({}), art(store, "AIDINQ.cbl"))


class TimesOutOnce(FakeClient):
    def create(self, **kw):
        if not getattr(self, "failed", False):
            self.failed = True
            raise TimeoutError("read timed out")
        return super().create(**kw)


def test_a_failure_during_a_run_is_recorded_and_the_rest_continue(store):
    a, b = art(store, "AIDINQ.cbl"), art(store, "AIDDBD.asm")
    client = TimesOutOnce({"purpose": "DBD.", "facts": [fact("Defines segment DISTRICT", [4, 4], "SEGM  NAME=DISTRICT", cat="data_read")]},
                          {"verdicts": [{"id": 1, "verdict": "supported"}]})
    res = DD.run(store, client, artifact_ids=[a["id"], b["id"]], force=True)
    assert res["analysed"] == 1 and len(res["errors"]) == 1 and "TimeoutError" in res["errors"][0]
    assert any(r["step"] == "deepdive" and not r["ok"] for r in store.runs())


def test_max_tokens_truncation_is_logged_as_a_failure(store):
    class Truncated(FakeClient):
        def create(self, **kw):
            m = super().create(**kw)
            m.stop_reason = "max_tokens"
            return m
    with pytest.raises(ValueError, match="truncated"):
        DD.analyse_file(store, Truncated(INQ, {"verdicts": []}), art(store, "AIDINQ.cbl"))
    assert any(r["step"] == "deepdive" and not r["ok"] and "truncated" in (r["error"] or "") for r in store.runs())


# ── invalidation and one job at a time ─────────────────────────────────────────────────────────────────────────

def test_changed_file_invalidates_its_analysis_and_deleted_files_are_pruned(store, monkeypatch):
    monkeypatch.setattr(DD, '_has_open_problem', lambda *a: False)
    a = art(store, "AIDINQ.cbl")
    DD.run(store, FakeClient(INQ, {"verdicts": []}), artifact_ids=[a["id"]], force=True)
    assert a["name"] not in {x["name"] for x in DD.pending(store)}
    store._db.execute("UPDATE artifact SET transcription = transcription || '\n* changed' WHERE id = ?", (a["id"],))
    assert a["name"] in {x["name"] for x in DD.pending(store)}
    dd = store.get_meta("deepdive")
    dd["99999"] = {"name": "gone.cbl", "facts": []}
    store.set_meta("deepdive", dd)
    DD.run(store, FakeClient({"purpose": "", "facts": []}), artifact_ids=[a["id"]], force=True)
    assert "99999" not in store.get_meta("deepdive")


def test_screens_and_documents_are_not_sent_for_code_review(store):
    assert not any(x["artifact_type"] == "ui_screen" for x in DD.pending(store))


def test_only_one_job_runs_per_program(tmp_path, monkeypatch):
    pytest.importorskip("fastapi")
    monkeypatch.setenv("CODESNAP_PROGRAMS", str(tmp_path / "programs"))
    from webapp import server
    need_fixture(RUN1)
    shutil.copytree(RUN1, tmp_path / "programs" / "live-test-09-29")
    import threading
    gate = threading.Event()

    class Slow(FakeClient):
        def create(self, **kw):
            gate.wait(5)
            return super().create(**kw)
    monkeypatch.setattr(server, "_client", lambda: Slow({"purpose": "", "facts": []}))
    server._DEEP.pop("live-test-09-29", None)
    j1 = server._deep_start("live-test-09-29")
    j2 = server._deep_start("live-test-09-29")
    assert j1 is j2 and j1["running"]
    gate.set()


# ── every rescan trigger, and a clean capture ──────────────────────────────────────────────────────────────────

def _art(text, **kw):
    return {"id": 1, "name": kw.pop("name", "x.cbl"), "transcription": text, "artifact_type": "code", "language": "COBOL", **kw}


class Store:
    def __init__(self, v=None):
        self.v = v

    def verification(self, _):
        return self.v

    def get_meta(self, _):
        return {}


def _flags(kind, line=3):
    from core.verify import REASONS
    return {"lines": 9, "verified": 8, "flags": [{"line": line, "text": "x", "reason": REASONS[kind]}]}


GOOD_COBOL = "\n".join(["       IDENTIFICATION DIVISION.", "       PROGRAM-ID. X.", "       PROCEDURE DIVISION.",
                        "           DISPLAY 'HI'.", "           STOP RUN."])


@pytest.mark.parametrize("kind, advice", [("cut", "scroll right"), ("break", "scroll back"), ("mismatch", "zoom in")])
def test_each_line_check_problem_asks_for_the_right_rescan(kind, advice):
    q = DD.capture_quality(Store(_flags(kind)), _art(GOOD_COBOL))
    assert q["status"] == "rescan" and q["issues"][0]["kind"] == kind and advice in q["issues"][0]["advice"]
    assert q["bad_lines"] == [3]


def test_cut_off_marker_in_the_text_asks_for_a_rescan():
    q = DD.capture_quality(Store({"lines": 5, "flags": []}), _art(GOOD_COBOL.replace("'HI'.", "'HI'. [CUT OFF]")))
    assert q["status"] == "rescan" and q["issues"][0]["kind"] == "cut"


def test_incomplete_file_asks_for_a_rescan():
    q = DD.capture_quality(Store({"lines": 2, "flags": []}), _art("       IDENTIFICATION DIVISION.\n       PROGRAM-ID. X."))
    assert q["status"] == "rescan" and any(i["kind"] == "partial" for i in q["issues"])


def test_compile_errors_are_not_reading_problems():
    q = DD.capture_quality(Store({"lines": 5, "flags": []}),
                           _art(GOOD_COBOL, validation_ok=0, validation_errors="x.cbl:4: syntax error near DISPLAY"))
    assert q["status"] != "rescan" and not any(i["kind"] == "compile" for i in q["issues"])


def test_model_flagged_lines_ask_for_a_rescan():
    q = DD.capture_quality(Store({"lines": 5, "flags": []}), _art(GOOD_COBOL), [{"line": 4, "reason": "garbled"}])
    assert q["status"] == "rescan" and 4 in q["bad_lines"]


def test_a_clean_verified_capture_is_good_and_an_unchecked_one_is_unchecked():
    assert DD.capture_quality(Store({"lines": 5, "verified": 5, "flags": []}), _art(GOOD_COBOL))["status"] == "good"
    assert DD.capture_quality(Store(None), _art(GOOD_COBOL))["status"] == "unchecked"


def test_report_wording_never_mentions_how_the_code_was_read():
    for kind in DD.REPORT_REASON:
        assert not re.search(r"scan|captur|screen|screenshot", DD.report_reason({"kind": kind, "reason": ""}), re.I)


# ── the report ─────────────────────────────────────────────────────────────────────────────────────────────────

def test_report_section_cites_every_statement_and_lists_the_rescan_once(store):
    import io
    import docx
    from core.report import docx_bytes
    a = art(store, "AIDINQ.cbl")
    DD.run(store, FakeClient(INQ, {"verdicts": [{"id": 1, "verdict": "supported"}, {"id": 2, "verdict": "supported"}]}),
           artifact_ids=[a["id"]], force=True)
    d = docx.Document(io.BytesIO(docx_bytes(store, rescan=True)))
    paras = [p.text for p in d.paragraphs]
    cells = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    # The report follows the template, so the per-component review section is not part of it.
    assert not any(p.startswith(("3.6 ", "Notes for ", "Appendix H", "Appendix I")) for p in paras)
    assert not re.search(r"\b(rescan|screenshots?|captur\w*|scann\w*)\b", "\n".join(paras), re.I)
    assert cells.count("Re-obtain a complete copy of AIDPAYRN.cbl") == 1


# ── the live eval's scoring (no API) ───────────────────────────────────────────────────────────────────────────

def test_eval_expectations_compile_and_score_as_intended():
    import importlib.util
    spec = importlib.util.spec_from_file_location("ddeval", Path(__file__).resolve().parent / "evals" / "deepdive_eval.py")
    ev = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ev)
    for groups in [g for f in ev.EXPECTED.values() for g in f.values() if not isinstance(g, str)] + list(ev.PLANTED[2].values()):
        for g in groups:
            re.compile(g)
    res = {"facts": [{"statement": "Holdback is 10% of gross (WS-HOLDBACK-PCT VALUE .100)", "quote": "VALUE .100", "lines": [70, 70]},
                     {"statement": "It writes a trailer record is written to PAYMENT-FILE", "quote": "", "lines": [1, 1]}]}
    sc = ev.score("AIDPAYRN.cbl", res)
    held = next(k for k in sc["found"] if k.startswith("pay period"))
    assert sc["found"]["holdback 10%"] and sc["found"][held]                 # not stated, so correctly held back
    assert not ev.score("AIDPAYRN.cbl", {"facts": [{"statement": "Pay period is 202609", "lines": [71, 71]}]})["found"][held]
    assert sc["traps"]
    for name in ev.EXPECTED:
        need_fixture(RUN1)
        assert (RUN1 / "sources" / f"{name}.v1").exists(), name      # every expectation is about a real fixture file


# ── the second look at lines nothing covers ────────────────────────────────────────────────────────────────────

def test_uncovered_finds_runs_of_code_with_no_findings_and_skips_comments():
    code = "\n".join(["# header", "a = 1", "b = 2", "c = 3", "", "d = 4", "e = 5", "f = 6", "g = 7", "}"])
    facts = [{"lines": [2, 3]}]
    assert DD.uncovered(code, facts) == ["6–9"]
    assert DD.uncovered(code, facts + [{"lines": [6, 9]}]) == []


def test_gap_pass_adds_what_the_first_pass_missed_without_duplicates(store):
    a = art(store, "AIDINQ.cbl")
    first = {"purpose": "x", "facts": [fact("Receives map AIDMAP1", [18, 18], "RECEIVE MAP('AIDMAP1')", cat="ui")]}
    gaps = {"facts": [fact("Receives map AIDMAP1", [18, 18], "RECEIVE MAP('AIDMAP1')", cat="ui"),
                      fact("Selects TOTAL_AMOUNT", [35, 35], "SELECT TOTAL_AMOUNT INTO :WS-TOTAL", cat="data_read")]}
    client = FakeClient(first, {"verdicts": []}, gaps=gaps)
    res = DD.analyse_file(store, client, a)
    assert [f["statement"] for f in res["facts"]] == ["Receives map AIDMAP1", "Selects TOTAL_AMOUNT"]
    assert "have NO findings yet" in client.calls[1]["messages"][0]["content"]
    assert any(r["step"] == "deepdive_gaps" for r in store.runs())


def test_capture_concerns_on_blank_or_missing_lines_are_ignored(store):
    a = art(store, "AidPaymentController.cs")
    blank = next(i for i, l in enumerate(a["transcription"].split("\n"), 1) if not l.strip())
    res = DD.analyse_file(store, FakeClient({"purpose": "", "facts": [], "capture_concerns": [
        {"line": blank, "reason": "trailing newline"}, {"line": 9999, "reason": "x"}, {"reason": "no line"}]}), a)
    assert res["capture_concerns"] == [] and res["quality"]["status"] != "rescan"


def test_column_based_languages_are_never_auto_rewritten():
    import team
    assert team._fixed_format("COBOL", "cbl") and team._fixed_format("", ".rpgle") and team._fixed_format("PL/I")
    assert not team._fixed_format("Python", "py")
