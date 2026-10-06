import io
import re
from types import SimpleNamespace

from docx import Document
from core.report import editorial


def test_colour_requires_evidence_and_never_confuses_missing_with_green():
    a = {"components": [{"name": "a"}], "verdict": {"code": "retain"},
         "confidence": {"level": "high"}, "total_risk": {"level": "low"}}
    assert editorial.decision(a, {"blockers": []})[0] == "Green"
    assert editorial.decision(a, {"blockers": ["review pending"]})[0] == "Yellow"
    a["verdict"]["code"] = "replace"
    assert editorial.decision(a, {})[0] == "Red"
    assert editorial.decision({}, {})[0] == "Not rated"


def test_report_keeps_only_template_sections_and_caps_prose():
    d = Document()
    d.add_heading("1. Application summary", 1)
    d.add_heading("1.1 Executive snapshot", 2)
    d.add_paragraph("Snapshot.")
    d.add_heading("1.2 Key findings", 2)
    for i in range(7):
        d.add_paragraph(f"Key finding number {i} is stated here in a short sentence.")
    d.add_heading("1.3 Immediate actions", 2)
    d.add_heading("2. Business context", 1)
    d.add_heading("2.1 Business purpose and ownership", 2)
    original = " ".join(f"Finding {i} needs a check before changes are approved." for i in range(25))
    d.add_paragraph(original)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Attribute"; table.cell(0, 1).text = "Entry"
    table.cell(1, 0).text = "Evidence"; table.cell(1, 1).text = "a.cbl:1"
    d.add_heading("3.6 Code analysis by component", 2)
    d.add_paragraph("One component has a detailed review.")
    d.add_heading("Appendix H. Extra", 1)
    d.add_paragraph("Extra material.")
    a = {"components": [{"name": "a.cbl", "language": "COBOL"}], "verdict": {"code": "retain", "label": "Retain"},
         "confidence": {"level": "low"}, "roadmap": {"phases": []}}
    meta = {}
    editorial.apply(SimpleNamespace(d=d), a, {"blockers": ["missing evidence"]}, {}, metadata=meta)
    assert meta["traffic_light"] == "yellow"
    paragraphs = d.paragraphs
    heads = [p.text for p in paragraphs if p.style.name.startswith("Heading")]
    assert "3.6 Code analysis by component" not in heads and not any(h.startswith("Appendix") for h in heads)
    assert "1.2 Key findings" in heads and "1.3 Immediate actions" in heads
    i = next(i for i, p in enumerate(paragraphs) if p.text == "1.2 Key findings")
    status = paragraphs[i + 1]
    assert status.text.startswith("Needs changes") and str(status.runs[0].font.color.rgb) == "996300"
    j = next(i for i, p in enumerate(paragraphs) if p.text == "1.3 Immediate actions")
    assert len([p for p in paragraphs[i + 1:j] if p.text.strip()]) <= 5
    section = next(i for i, p in enumerate(paragraphs) if p.text == "2.1 Business purpose and ownership")
    assert len(editorial.lines(paragraphs[section + 1].text)) <= 10
    from core.report.html import render_document
    out = io.BytesIO(); d.save(out)
    assert 'color:#996300;font-weight:700' in render_document(out.getvalue(), {})
    assert d.tables[0].cell(0, 0).text == "Attribute"   # the template's column names are kept


def test_numbered_actions_remain_whole_and_html_preserves_risk_colours():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from core.report.html import render_document
    d = Document()
    d.add_heading("12. Modernization options and recommendation", 1)
    d.add_heading("12.4 Sequencing and timing", 2)
    action = "1. Confirm " + "the required evidence " * 8 + ". Nothing is rebuilt before this decision."
    d.add_paragraph(action)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Impact 5"
    t.cell(0, 1).text = "25"
    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "F4C7C3")
    t.cell(0, 1)._tc.get_or_add_tcPr().append(shade)
    identifier = "LONG_SOURCE_REFERENCE_" * 30
    t.cell(1, 0).text = "Source"; t.cell(1, 1).text = identifier
    editorial.apply(SimpleNamespace(d=d), {}, {}, {})
    assert any(p.text == action for p in d.paragraphs)
    assert t.cell(1, 1).text == identifier   # a cell with no sentence break is left whole
    assert not any(p.text.startswith("Appendix") for p in d.paragraphs)
    data = io.BytesIO(); d.save(data)
    html = render_document(data.getvalue(), {})
    assert "background:#F4C7C3" in html
    assert "<th" in html
