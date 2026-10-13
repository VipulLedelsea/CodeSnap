"""Report layout: evidence tags become plain references, long registers are capped, columns follow their content,
wide tables get a landscape page that survives the cleanup pass, and open items are merged."""
import io

import pytest

docx = pytest.importorskip("docx")
from docx.oxml.ns import qn

from core import exec_summary as ES
from core.report import editorial
from core.report.template_docx import _merge_open_items

INDEX = {"A.cbl:F1": {"lines": [10, 40], "statement": "S1"}, "B.cbl:F2": {"lines": [5, 9], "statement": "S2"},
         "C.cbl:F3": {"lines": [1, 3], "statement": "S3"}}


def summary(**kw):
    base = {"verdict": "Remediate", "provisional": True, "what_it_does": "It pays people. [A.cbl:F1]",
            "reasons": ["r1", "r2", "r3"], "lifecycle": "not known", "lifecycle_verified": False,
            "decision": "Decide soon.", "confidence": "medium", "confidence_and_coverage": "All files read.",
            "needed_to_finalize": ["a", "b", "c", "d", "e", "f"],
            "risks": [{"text": "Low thing [C.cbl:F3]", "severity": "Low", "facts": ["C.cbl:F3"]},
                      {"text": "Big thing [A.cbl:F1] [B.cbl:F2]", "severity": "High", "facts": ["A.cbl:F1", "B.cbl:F2"]}]}
    base.update(kw)
    return base


def test_plain_refs_replace_tags_with_file_and_lines():
    assert ES.plain_refs("It pays. [A.cbl:F1] [B.cbl:F2]", INDEX) == "It pays. (A.cbl lines 10-40; B.cbl lines 5-9)"
    assert "[" not in ES.plain_refs("x [Z.cbl:F9]", INDEX)


def test_report_paragraphs_are_short_and_carry_no_tags():
    out = ES.report_paragraphs(summary(), INDEX)
    text = " ".join(out["before"] + out["risks"] + out["after"])
    assert "[" not in text and ":F" not in text
    assert out["risks"][0].startswith("High:")                     # most serious first
    assert "and 2 more (Section 13.4)" in out["before"][0]        # only four missing items named
    assert len(out["before"]) + len(out["after"]) <= 6


def test_evidence_rows_list_each_cited_finding_once():
    rows = ES.evidence_rows(summary(), INDEX)
    assert [r[0] for r in rows] == ["A.cbl", "B.cbl", "C.cbl"] and rows[0][1] == "10-40"


def test_open_items_merge_by_owner_and_section_group():
    items = [("Security control: authentication (Section 8.1)", "8.1", "Information security"),
             ("Security control: authorization (Section 8.1)", "8.1", "Information security"),
             ("Client policy (Section 8.5)", "8.5", "Information security"),
             ("Removed after the code check: not shown by the code.", "4.2", "IT"),
             ("Hosting", "3.1", "IT")]
    out = _merge_open_items(items)
    assert len(out) == 2
    assert "Security control: authentication, authorization" in out[0][0] and "Client policy" in out[0][0]
    assert out[1][0] == "Hosting"


def build_doc(rows=12, cols=9):
    d = docx.Document()
    d.add_paragraph("9.1 Risk register", "Heading 2")
    t = d.add_table(rows=rows + 1, cols=cols)
    for j in range(cols):
        t.cell(0, j).text = "Column %d" % j
    for i in range(1, rows + 1):
        for j in range(cols):
            t.cell(i, j).text = ("R-%02d" % i) if j == 0 else ("word " * 30) if j == 1 else "x"
    return d


def test_columns_follow_content():
    d = build_doc()
    widths = editorial._fit_columns(d.tables[0], int(d.sections[0].page_width - d.sections[0].left_margin - d.sections[0].right_margin))
    assert widths[1] == max(widths)            # the description column gets the room
    assert widths[0] < widths[1] / 2


def test_wide_text_heavy_table_gets_a_landscape_section_that_survives_cleanup():
    d = build_doc(rows=30)
    editorial.WIDE.clear()
    editorial._landscape_wide_tables(d)
    sects = list(d.element.body.iter(qn("w:sectPr")))
    assert len(sects) == 3 and sects[1].find(qn("w:pgSz")).get(qn("w:orient")) == "landscape"
    editorial._simple_english(d)               # removes empty paragraphs; the section breaks must stay
    assert len(list(d.element.body.iter(qn("w:sectPr")))) == 3


def test_small_wide_table_stays_portrait():
    d = build_doc(rows=5)
    editorial.WIDE.clear()
    editorial._landscape_wide_tables(d)
    assert len(list(d.element.body.iter(qn("w:sectPr")))) == 1


def test_register_over_its_cap_moves_rows_to_the_appendix_list():
    d = build_doc(rows=20, cols=8)
    out = editorial._cap_registers(d)
    assert len(d.tables[0].rows) == 1 + 6 and out[0][0] == "Risk register" and len(out[0][2]) == 14
