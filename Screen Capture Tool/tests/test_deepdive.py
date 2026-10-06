"""The deep code analysis is held to the evidence: quotes must be in the file, lines are corrected, unsupported
statements are dropped by the second review, statements resting on unreadable lines are held back, cross-file
observations may only cite surviving facts, and poor captures produce a rescan request."""
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest

from conftest import need_fixture

from core import deepdive as DD
from core.model import ProgramStore

RUN1 = Path(__file__).resolve().parent / "fixtures" / "run1_program"

CODE = """PROGRAM-ID. PAYCALC.
    MOVE ENT-BASE-AMOUNT TO WS-GROSS
    COMPUTE WS-HOLDBACK = WS-GROSS * WS-HOLDBACK-PCT
    SUBTRACT WS-HOLDBACK FROM WS-GROSS GIVING WS-NET
    WRITE PAYMENT-REC
    05  PAY-NET  PIC S9(11)V99 COMP-3. [CUT OFF]
    STOP RUN."""


def tool_msg(name, data):
    return SimpleNamespace(content=[SimpleNamespace(type="tool_use", name=name, input=data)],
                           usage=SimpleNamespace(input_tokens=100, output_tokens=50), stop_reason="tool_use")


class FakeClient:
    """Plays the model, answering by the tool each call asks for: analysis, then (if any) the gap pass, the review and
    the program synthesis. A tool asked for more often than it has answers gets an empty answer."""

    def __init__(self, analysis, review=None, program=None, gaps=None):
        self.answers = {"record_analysis": [analysis] + ([gaps] if gaps else []),
                        "record_review": [review] if review else [], "record_program": [program] if program else []}
        self.calls = []
        self.messages = self

    @property
    def queue(self):
        return [(k, v) for k, vs in self.answers.items() for v in vs]

    @queue.setter
    def queue(self, items):
        self.answers = {"record_analysis": [], "record_review": [], "record_program": []}
        for k, v in items:
            self.answers[k].append(v)

    def create(self, **kw):
        self.calls.append(kw)
        name = (kw.get("tools") or [{}])[0].get("name", "record_program")
        q = self.answers.get(name) or []
        data = q.pop(0) if q else ({"facts": []} if name == "record_analysis" else
                                   {"verdicts": []} if name == "record_review" else {"observations": []})
        return tool_msg(name, data)


def fact(stmt, lines, quote, cat="calculation", basis="observed", **kw):
    return {"category": cat, "statement": stmt, "lines": lines, "quote": quote, "basis": basis, **kw}


def test_quotes_must_be_in_the_file_and_lines_are_corrected():
    facts = [fact("Holdback is gross times the holdback percentage", [3, 3], "WS-GROSS * WS-HOLDBACK-PCT"),
             fact("Net is written to the payment record", [1, 1], "WRITE PAYMENT-REC", cat="data_write"),
             fact("A 2% late fee is added", [3, 3], "ADD WS-LATE-FEE"),
             fact("Probably runs monthly", [7, 7], "STOP RUN.", basis="inferred")]
    kept, corrected, rejected, unver = DD.check_facts(facts, CODE)
    assert [f["statement"][:8] for f in kept] == ["Holdback", "Net is w"]
    assert kept[1]["lines"] == [5, 5] and corrected == 1                       # moved to where the quote really is
    assert {r["why"] for r in rejected} == {"quote not found in the file", "inferred without reasoning"}
    assert not unver


def test_facts_resting_only_on_unreadable_lines_are_held_back():
    kept, _, _, unver = DD.check_facts([fact("Net is packed decimal", [6, 6], "PIC S9(11)V99 COMP-3", cat="data_write")],
                                       CODE, bad=[6])
    assert not kept and len(unver) == 1


def test_whitespace_differences_in_quotes_are_tolerated_but_not_wording():
    kept, _, rejected, _ = DD.check_facts([fact("x", [2, 2], "MOVE   ENT-BASE-AMOUNT  TO WS-GROSS"),
                                           fact("y", [2, 2], "MOVE ENT-BASE TO WS-GROSS")], CODE)
    assert len(kept) == 1 and len(rejected) == 1


@pytest.fixture
def store(tmp_path):
    need_fixture(RUN1)
    shutil.copytree(RUN1, tmp_path / "live-test-09-29")
    st = ProgramStore.open("live-test-09-29", root=tmp_path)
    yield st
    st.close()


def art(store, name):
    return next(a for a in store.artifacts() if a["name"] == name)


def test_cut_off_lines_produce_a_rescan_request_with_advice(store):
    q = DD.capture_quality(store, art(store, "AIDPAYRN.cbl"))
    assert q["status"] == "rescan" and q["bad_lines"]
    cut = next(i for i in q["issues"] if i["kind"] == "cut")
    assert cut["lines"] and "scroll right" in cut["advice"]
    names = [r["name"] for r in DD.rescan_requests(store)]
    assert "AIDPAYRN.cbl" in names and "AidPaymentController.cs" not in names


def test_model_capture_concerns_also_trigger_a_rescan(store):
    a = art(store, "AidPaymentController.cs")
    assert DD.capture_quality(store, a)["status"] != "rescan"
    q = DD.capture_quality(store, a, [{"line": 12, "reason": "string literal is not closed"}])
    assert q["status"] == "rescan" and 12 in q["bad_lines"]


def test_analyse_file_end_to_end_with_review(store):
    a = art(store, "AidPaymentController.cs")
    analysis = {"purpose": "An MVC controller for district aid (lines 8-43).",
                "facts": [fact("Approve updates dbo.PaymentBatch status to 'A'", [28, 28], "SET Status = 'A'", cat="data_write"),
                          fact("Approve has no [Authorize] attribute", [24, 25], "public ActionResult Approve", cat="security",
                               severity="high"),
                          fact("The connection string holds a password", [12, 12], "Password=P@ssw0rd1", cat="security",
                               severity="high"),
                          fact("Approve sends an email to the district", [26, 26], "SendMail(district)", cat="interface")],
                "unknowns": [{"what": "PaymentService is not provided", "line": 11}]}
    review = {"verdicts": [{"id": 1, "verdict": "supported"},
                           {"id": 2, "verdict": "partly", "corrected": "Approve has no [Authorize] attribute on the method or class"},
                           {"id": 3, "verdict": "unsupported", "why": "test"}]}
    client = FakeClient(analysis, review)
    res = DD.analyse_file(store, client, a)
    st = [f["statement"] for f in res["facts"]]
    assert st == ["Approve updates dbo.PaymentBatch status to 'A'",
                  "Approve has no [Authorize] attribute on the method or class"]
    assert res["rejected"] == 2                                  # invented email + one the review rejected
    assert res["reviewed"]["supported"] == 1 and res["reviewed"]["partly"] == 1
    assert client.calls[0]["thinking"]["type"] == "adaptive"     # the strongest model reasons before answering
    assert "Password=P@ssw0rd1" in client.calls[0]["messages"][0]["content"]
    assert store.runs()[-1]["step"] == "deepdive_review"


def test_pending_tracks_changed_files_and_run_stores_results(store):
    before = {a["name"] for a in DD.pending(store)}
    assert "AIDPAYRN.cbl" in before
    only = art(store, "AIDINQ.cbl")
    client = FakeClient({"purpose": "CICS inquiry.", "facts": [fact("Receives map AIDMAP1", [18, 18], "RECEIVE MAP('AIDMAP1')", cat="ui")]},
                        {"verdicts": [{"id": 1, "verdict": "supported"}]})
    res = DD.run(store, client, artifact_ids=[only["id"]])
    assert res["analysed"] >= 1 and str(only["id"]) in store.get_meta("deepdive")
    assert "AIDINQ.cbl" not in {a["name"] for a in DD.pending(store)}


def test_program_synthesis_may_only_cite_surviving_facts(store):
    first = store.add_artifact('a.cbl', transcription='\n\nMOVE 6728 TO RATE')
    second = store.add_artifact('b.rpg', transcription='\n' * 8 + 'EVAL TOTAL = ADM * RATE')
    store.set_meta("deepdive", {str(first): {"name": "a.cbl", "hash": DD._hash(store.artifact(first)['transcription']), "facts": [{"id": 1, "category": "calculation", "statement": "rate 6728",
                                                                  "lines": [3, 3]}]},
                                str(second): {"name": "b.rpg", "hash": DD._hash(store.artifact(second)['transcription']), "facts": [{"id": 1, "category": "calculation", "statement": "ADM * RATE",
                                                                  "lines": [9, 9]}]}})
    client = FakeClient({}, None, None)
    client.queue = [("record_program", {"observations": [
        {"title": "Rate held twice", "statement": "The rate is held in two places", "facts": [f"F{first}.1", f"F{second}.1"]},
        {"title": "Invented", "statement": "Both write to the ledger", "facts": ["F9.9"]}]})]
    out = DD.synthesize(store, client)
    assert [o["title"] for o in out["observations"]] == ["Rate held twice"]
    assert out["observations"][0]["cites"] == ["a.cbl lines 3–3", "b.rpg lines 9–9"]
    store.fill_artifact(first, artifact_type='code', language='COBOL', transcription='\n\nMOVE 7000 TO RATE')
    calls = len(client.calls)
    assert DD.synthesize(store, client)['observations'] == []
    assert len(client.calls) == calls  # stale reviews cannot trigger synthesis or provider spend


def test_report_shows_the_review_with_line_citations_and_the_rescan(store):
    import io
    import docx
    from core.report import docx_bytes
    a = art(store, "AidPaymentController.cs")
    store.set_meta("deepdive", {str(a["id"]): {"name": a["name"], "hash": DD._hash(a["transcription"]), "prompt_version": DD.PROMPT_VERSION,
                                               "ran_at": "2026-09-30T00:00:00+00:00", "purpose": "Controller for aid payments.",
                                               "facts": [{"id": 1, "category": "data_write", "basis": "observed", "lines": [28, 28],
                                                          "statement": "Approve sets dbo.PaymentBatch.Status to 'A'"}],
                                               "rejected": 1, "unverifiable": [], "unknowns": []}})
    d = docx.Document(io.BytesIO(docx_bytes(store, rescan=True)))
    text = "\n".join(p.text for p in d.paragraphs)
    cells = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "Approve sets dbo.PaymentBatch.Status to 'A' (line 28)." in text
    assert "1 were rejected in these checks" in text and "Not yet reviewed line by line." in text
    assert "Blocks issue: Re-obtain a complete copy of AIDPAYRN.cbl" in cells and "copy incomplete (lines cut short" in cells
    assert "A complete copy is needed" in text
    assert not __import__("re").search(r"\b(rescan|screenshot|captur\w*)\b", text + cells, __import__("re").I)
    assert 'full count unconfirmed' in cells


def test_app_endpoints_show_quality_and_run_the_analysis(tmp_path, monkeypatch):
    pytest.importorskip("fastapi")
    import time as _t
    from fastapi.testclient import TestClient
    monkeypatch.setenv("CODESNAP_PROGRAMS", str(tmp_path / "programs"))
    from webapp import server
    need_fixture(RUN1)
    shutil.copytree(RUN1, tmp_path / "programs" / "live-test-09-29")
    a_id = None
    with ProgramStore.open("live-test-09-29") as st:
        a_id = art(st, "AIDINQ.cbl")["id"]
    fake = FakeClient({"purpose": "Inquiry.", "facts": [fact("Receives map AIDMAP1", [18, 18], "RECEIVE MAP('AIDMAP1')", cat="ui")]},
                      {"verdicts": [{"id": 1, "verdict": "supported"}]})
    monkeypatch.setattr(server, "_client", lambda: fake)
    monkeypatch.setattr(server, "api_key_status", lambda: {"has_key": True})
    api = TestClient(server.app)
    d = api.get("/api/programs/live-test-09-29").json()
    pay = next(x for x in d["artifacts"] if x["name"] == "AIDPAYRN.cbl")
    assert pay["quality"]["status"] == "rescan"
    r = api.post(f"/api/programs/live-test-09-29/deepdive?artifact_id={a_id}").json()
    assert r["ok"]
    deadline = _t.monotonic() + 30          # poll on a deadline, not a fixed number of tries, so a slow machine is not a failure
    while True:
        s = api.get("/api/programs/live-test-09-29/deepdive").json()
        if not s["job"]["running"] or _t.monotonic() > deadline:
            break
        _t.sleep(0.05)
    assert not s["job"]["running"], "the analysis job did not finish in 30 seconds"
    assert str(a_id) in s["files"] and any(x["name"] == "AIDPAYRN.cbl" for x in s["rescan"])
    detail = api.get(f"/api/programs/live-test-09-29/artifacts/{a_id}").json()
    assert detail["deep"]["facts"][0]["statement"] == "Receives map AIDMAP1" and detail["deep_current"]
    rep = api.get("/api/programs/live-test-09-29/report.docx")
    assert rep.status_code in (200, 409)
