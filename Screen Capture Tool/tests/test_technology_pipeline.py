"""Offline plumbing tests, not a substitute for live model/dialect accuracy."""
import pytest
from core.model import ProgramStore, ingest_capture, ingest_artifact
from core.technology_support import REGISTRY, resolve, analysis_context, program_coverage, stack_category
from core.transcription_formats import TECHNOLOGIES
from test_ingest import FakeClient
from test_multiformat_spacing import SOURCES
from render import render_code

@pytest.mark.parametrize('technology',TECHNOLOGIES)
def test_every_requested_label_has_source_analysis_guidance(technology):
    scope=resolve(language=technology)
    assert scope['technology']==technology
    assert scope['visible_only']
    assert technology in analysis_context(language=technology)
    assert 'Screens are the only evidence' in analysis_context(language=technology)

@pytest.mark.parametrize('filename,language,expected',[
    ('account.kt','','Kotlin'),('account.fs','','F#'),('account.tsx','','React'),
    ('account.rs','','Rust'),('report.rdl','','SSRS'),('model.dtsx','','SSIS'),
    ('legacy.cls','VB6','VB6'),('account.m','Objective-C','Objective-C'),
    ('app.src','Java','Java SE'),('app.src','PL/I','PL/I'),
])
def test_explicit_language_resolves_before_ambiguous_extension(filename,language,expected):
    assert resolve(filename,language)['technology']==expected

@pytest.mark.parametrize('name',SOURCES)
def test_capture_source_reaches_analysis_and_report_inventory(name,tmp_path):
    # Mock structure output verifies the pipeline, not the provider's reasoning.
    source=SOURCES[name]
    image=render_code(source,tmp_path/'capture.png')
    client=FakeClient([{'entities':[],'relations':[]}])
    with ProgramStore.create('Screen source pipeline',root=tmp_path/'programs') as store:
        aid=ingest_capture(store,client,[image],{'code':source,'code_name':'visible_source',
            'extension':'src','language':name,'is_code':True,'errors':'Not verified: no compiler check'})
        artifact=store.artifact(aid)
        assert artifact['transcription']==source
        assert artifact['status']=='structured'
        assert artifact['validation_ok'] is None
        assert len(store.artifact_evidence(aid))==1
        rows=program_coverage(store)
        assert rows[0]['file']=='visible_source.src'
        assert rows[0]['method']!='Analysis method not recorded'
        assert 'hidden source' in rows[0]['scope']
        from core.report import build
        report=build(store)
        section=next(s for s in report['sections'] if s['id']=='coverage')
        table=next(b for b in section['blocks'] if b.get('head')==['Component','Source format','Analysis method','Status','Scope'])
        assert table['rows'][0][0]=='visible_source.src'
        assert table['rows'][0][3]=='structured'


def test_visual_config_gets_model_analysis_instead_of_generic_inventory(tmp_path):
    with ProgramStore.create('Visible SSIS',root=tmp_path) as store:
        source='<Package><Task name="LoadAccounts" /></Package>'
        aid=store.add_artifact('visible.dtsx','config','SSIS',source)
        client=FakeClient([{'entities':[{'kind':'job','name':'LoadAccounts','line_start':1,'line_end':1}], 'relations':[]}])
        ingest_artifact(store,client,aid)
        assert client.calls
        assert store.entity_by_key('job:LoadAccounts')
        assert not program_coverage(store)[0]['limited']
        assert 'SSIS' in client.calls[0]['messages'][0]['content']


def test_visual_config_offline_inventory_is_marked_limited(tmp_path):
    with ProgramStore.create('Visible SSIS',root=tmp_path) as store:
        aid=store.add_artifact('visible.dtsx','config','SSIS','<Package><Task name="LoadAccounts" /></Package>')
        ingest_artifact(store,None,aid)
        assert program_coverage(store)[0]['limited']


def test_model_failure_does_not_silently_mark_unknown_language_complete(tmp_path):
    with ProgramStore.create('Failure',root=tmp_path) as store:
        aid=store.add_artifact('main.kt','code','Kotlin','fun main() { println(12) }')
        with pytest.raises(RuntimeError,match='api down'):
            ingest_artifact(store,FakeClient(fail=True),aid)
        assert store.artifact(aid)['status']=='failed'


def test_truncated_analysis_cannot_be_published_as_complete(tmp_path):
    with ProgramStore.create('Truncated',root=tmp_path) as store:
        aid=store.add_artifact('main.kt','code','Kotlin','fun main() { println(12) }')
        with pytest.raises(ValueError,match='truncated'):
            ingest_artifact(store,FakeClient(stop_reason='max_tokens'),aid)
        assert store.artifact(aid)['status']=='failed'
        assert store.usage()['failures']==1


def test_report_separates_products_from_languages():
    assert stack_category('Kotlin')=='Programming languages'
    assert stack_category('React')=='Application frameworks and platforms'
    assert stack_category('SSIS')=='Development and data tooling'
    assert stack_category('CICS')=='Integration and middleware'

@pytest.mark.parametrize('name',SOURCES)
def test_new_source_formats_reach_quote_checked_independent_review(name,tmp_path):
    from core import deepdive as DD
    from test_deepdive import FakeClient as ReviewClient, fact
    source=SOURCES[name]
    row,quote=max(enumerate(source.splitlines(),1),key=lambda item:len(item[1]))
    quote=quote.strip()
    with ProgramStore.create('Visible review',root=tmp_path) as store:
        aid=store.add_artifact('visible.src','code',name,source)
        store.set_status(aid,'structured')
        client=ReviewClient({'purpose':'Visible source review','facts':[
            fact('The quoted source is present',[row,row],quote,cat='control_flow'),
            fact('An absent operation exists',[row,row],'NOT_PRESENT_OPERATION()',cat='control_flow')]},
            {'verdicts':[{'id':1,'verdict':'supported'}]})
        result=DD.analyse_file(store,client,store.artifact(aid))
        assert len(result['facts'])==1
        assert result['facts'][0]['quote']==quote
        assert result['facts'][0]['review']=='supported'
        assert result['rejected']>=1
        context=analysis_context('visible.src',name)
        if context:
            assert context in client.calls[0]['messages'][0]['content']
            review_call=next(call for call in client.calls if call['tools'][0]['name']=='record_review')
            assert context in review_call['messages'][0]['content']
        store.set_meta('deepdive',{str(aid):result})
        assert program_coverage(store)[0]['review_status']=='Detailed evidence review recorded'


def test_truncated_forensic_analysis_remains_incomplete(tmp_path):
    from core import deepdive as DD
    from test_deepdive import FakeClient as ReviewClient
    class Truncated(ReviewClient):
        def create(self,**kwargs):
            message=super().create(**kwargs)
            message.stop_reason='max_tokens'
            return message
    with ProgramStore.create('Incomplete',root=tmp_path) as store:
        aid=store.add_artifact('main.kt','code','Kotlin','fun main() { println(12) }')
        with pytest.raises(ValueError,match='truncated'):
            DD.analyse_file(store,Truncated({'facts':[]}),store.artifact(aid))
        assert program_coverage(store)[0]['review_status']=='Detailed evidence review pending'


def test_template_technology_table_retains_framework_and_tool_labels(tmp_path):
    import io
    from docx import Document
    from core.report import docx_bytes
    with ProgramStore.create('Technology table',root=tmp_path) as store:
        for name,language,kind,source in [
            ('Account.tsx','React','web','export function Account() { return <span>12</span>; }'),
            ('Package.dtsx','SSIS','config','<Package><Task name="LoadAccounts" /></Package>')]:
            aid=store.add_artifact(name,kind,language,source)
            ingest_artifact(store,FakeClient(),aid)
        doc=Document(io.BytesIO(docx_bytes(store,pdf=False)))
        stack_table=next(table for table in doc.tables if any('Technology and version' in cell.text for cell in table.rows[0].cells))
        stack='\n'.join(cell.text for row in stack_table.rows for cell in row.cells)
        assert 'React (source label; deployed use and version not confirmed)' in stack
        assert 'SSIS (source label; deployed use and version not confirmed)' in stack
        alltext='\n'.join(p.text for p in doc.paragraphs)
        assert 'same business data is also copied across platforms' not in alltext



def test_context_diagram_retains_components_on_unclassified_platforms():
    from core.report.figures import context
    scene=context({'components':[{'name':'Account.kt','platform':'Other','layer':'Application','role':'Program'}]},
                  {'occurrences':[]},'A long application name for context validation')
    assert any(node['title']=='Account.kt' for node in scene['nodes'])
    assert any(group['title']=='Other' for group in scene['groups'])
    from core.diagrams.layout import text_w
    for group in scene['groups']:
        assert group['w']>=text_w(group['title'],True)+20
