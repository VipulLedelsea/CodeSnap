"""Guard technical prose against claims that source-only review cannot establish."""
import io
import shutil
from pathlib import Path

import docx
import pytest

from conftest import need_fixture

from core.model import ProgramStore
from core.report import docx_bytes
from core.report import architecture as A
from core.report import controls as C
from core.report import dataarch as DA
from core.report import wording as W
from core.report.template_docx import _infer_stack, _sor


class EmptyStore:
    def entities(self, *args):
        return []

    def findings(self):
        return []

    def coverage(self):
        return {}

    def artifacts(self):
        return []


@pytest.fixture(scope="module")
def prose(tmp_path_factory):
    root = tmp_path_factory.mktemp("grounding")
    fixture = Path(__file__).parent / "fixtures/demo930"
    need_fixture(fixture)
    shutil.copytree(fixture, root / "demo")
    with ProgramStore.open("demo", root=root) as store:
        document = docx.Document(io.BytesIO(docx_bytes(store, rescan=True)))
    return "\n".join([p.text for p in document.paragraphs] +
                     [c.text for table in document.tables for row in table.rows for c in row.cells])


def test_missing_tests_do_not_establish_zero_coverage(prose):
    assert "Not measured; no test files identified in the supplied source" in prose
    assert "0% (no automated tests" not in prose
    assert "Tests were never built" not in prose


def test_file_complexity_is_not_presented_as_measured_routine_complexity(prose):
    assert "routine-level complexity not measured" in prose
    assert "Under 15 per routine" not in prose


def test_report_does_not_assert_offline_staffing_or_guaranteed_outcomes(prose):
    for text in ["small, shrinking pool", "errors reach users", "any user of the PC can change",
                 "outside business approval and audit", "public money", "in daily use", "should changing",
                 "no change can be verified automatically"]:
        assert text not in prose


def test_platform_matching_does_not_confuse_javascript_with_java():
    assert A._platform("JavaScript", "ui.js") == "Web browser"
    assert A._platform("Java", "Service.java") == "Java"
    assert A._platform("C#", "Service.cs") == "Microsoft .NET (host to confirm)"


def test_language_alone_does_not_establish_database_or_web_host():
    stack = {}
    _infer_stack(stack, [{"name": "Service.cs", "language": "C#", "transcription": "class Service {}"}],
                 {"sql_server": [], "db2": False}, {"platforms": ["Microsoft .NET (host to confirm)"]}, EmptyStore())
    text = str(stack)
    assert "Windows Server" not in text and "IIS" not in text
    assert "to confirm" in text
    from core.report.options import components
    component = {"name": "Service.cs", "language": "C#", "platform": "Microsoft .NET (host to confirm)",
                 "layer": "Application", "role": "Program", "writes": [], "lines": 1, "routines": []}
    model = {"components": [component]}
    assert components(model, [], [], {"Service.cs": "class Service {}"})[0]["code"] == "retain"
    assert components(model, [], [], {"Service.cs": "using System.Web.Mvc;"})[0]["code"] == "replatform"
    stack = {}
    _infer_stack(stack, [{"name": "run.cbl", "language": "COBOL", "transcription": "EXEC SQL SELECT 1 END-EXEC"}],
                 {"sql_server": [], "db2": True}, {"platforms": []}, EmptyStore())
    assert "Db2 for z/OS" not in str(stack)
    assert "database engine and version not confirmed" in str(stack)


def test_multiple_writers_do_not_establish_owner_or_missing_documentation():
    comp = {"name": "a", "platform": "Python runtime", "layer": "Application", "routines": ["CALCULATE"]}
    observations = A.observations(EmptyStore(), {}, [comp, {**comp, "name": "b"}],
                                 {"ledger": {"readers": set(), "writers": {"a", "b"}}})
    text = str(observations)
    assert "no single owner" not in text and "undocumented outside the code" not in text
    assert "Confirm ownership" in text and "Routine names suggest" in text


def test_read_only_does_not_establish_system_of_record():
    class Store:
        def relations(self, **kwargs):
            return [{"kind": "reads"}]

    assert _sor(Store(), {"id": 1, "attrs": {}}) == "To confirm (read only in the supplied code)"


def test_trailer_identifier_is_not_credited_as_a_written_control_record():
    dm = {"occurrences": [{"business": "Payment", "engine": "IBM mainframe (files)", "access": "W",
                           "component": "pay.cbl", "store": "PAYMENT-FILE"}], "entities": []}
    controls = C.assess(EmptyStore(), [{"name": "pay.cbl", "transcription": "01 TRAILER PIC X(80)."}],
                        [], dm, [], {"tests": 0}, [], [], True)
    row = next(r for r in controls if r[0] == "Control totals on file exchanges")
    assert row[3] is None and "to confirm" in row[1]
    assert "a trailer or control record is written" not in row[1]
    assert "implementation to confirm" in str(DA.reconciliation([{"name": "pay.cbl", "transcription": "01 TRAILER PIC X."}]))


def test_file_names_and_evidence_survive_editorial_cleanup():
    assert W.clean("3 file(s): capture.py and scan.cs", verbatim=["capture.py", "scan.cs"]) == "3 files: capture.py and scan.cs"
    evidence = {"snippet": 'var capture = "student data";', "quote": "scanned", "source_line": "capture();"}
    assert W.scrub(evidence) == evidence
    assert W.scrub({"name": "capture.py"}, verbatim=["capture.py"]) == {"name": "capture.py"}


def test_report_cache_rebuilds_after_renderer_revision_changes(tmp_path, monkeypatch):
    from core import report as R
    calls = []

    def package(store, **kwargs):
        calls.append(True)
        out = store.exports_dir / "report"
        out.mkdir(parents=True, exist_ok=True)
        slug = store.info["slug"]
        for suffix in ["assessment_report.html", "assessment_report.docx", "diagrams.vsdx", "diagrams.drawio"]:
            (out / f"{slug}_{suffix}").write_bytes(b"fixture")
        return {"zip": b"fixture", "verdict": {}}

    monkeypatch.setattr(R, "package", package)
    with ProgramStore.create("Cache", root=tmp_path) as store:
        assert R.cached_package(store)["built"] is True
        assert R.cached_package(store)["built"] is False
        monkeypatch.setattr(R, "REPORT_REVISION", "changed-renderer")
        assert R.cached_package(store)["built"] is True
        assert len(calls) == 2


def test_introductions_do_not_claim_complete_inventory_or_measured_operations(prose):
    assert "inventories every technical component" not in prose
    assert "complete inventory of the integrations" not in prose
    assert "quantifies the effort required to keep" not in prose
    assert "assessment judgments, not measured incident probabilities" in " ".join(prose.split())


@pytest.mark.parametrize("attrs, label", [
    ({"hierarchy": "child of"}, "child of"),
    ({"logical": True}, "logical relationship"),
    ({"foreign_key": ["district_id"]}, "FK"),
    ({"trigger_on": True}, "trigger"),
    ({}, "dependency")
])
def test_data_diagram_labels_only_evidenced_relationships(attrs, label):
    from core.diagrams.build import _data_relationship_label
    assert _data_relationship_label(attrs) == label


def test_ims_nodes_are_segments_in_the_data_diagram():
    from core.diagrams.build import _data_stereotype
    assert _data_stereotype({"origin": "source", "attrs": {"ims_segment": True}}) == "IMS segment"
    assert _data_stereotype({"origin": "source", "attrs": {}}) == "table"


@pytest.mark.parametrize('language,text,expected', [
    ('COBOL', 'EXEC SQL SELECT * FROM PAYMENT END-EXEC', 'Database (engine not identified)'),
    ('RPG', 'dcl-f PAYMENT usage(*input);', 'IBM i (database engine to confirm)'),
    ('Java', 'jdbc:db2://host:50000/database', 'IBM Db2 (host to confirm)'),
])
def test_storage_engine_requires_explicit_database_evidence(language, text, expected):
    artifact = {'language': language, 'transcription': text}
    entity = {'kind': 'table', 'name': 'PAYMENT', 'attrs': {}}
    assert DA._engine(artifact, entity) == expected


def test_report_does_not_assume_schedule_or_external_documentation(prose):
    assert 'Monthly payment runs' not in prose
    assert 'file layouts documented only in code' not in prose
    assert 'no automated tests.' not in prose


def test_comment_years_are_not_calculation_multipliers():
    arts = [{'name': 'job.jcl', 'transcription': '//* 1996 nightly intake\n//S EXEC PGM=IEFBR14'},
            {'name': 'layout.cbl', 'transcription': '      * 1996 record\n       01 RECORD.'}]
    assert C.hard_rates(arts) == []
    assert C.hard_rates([{'name':'x.cbl','transcription':'           COMPUTE TOTAL = AMOUNT * 1250'}]) == ['multiplier 1250 in x.cbl']


def test_cobol_roles_follow_source_not_language_alone():
    copy = {'name': 'layout.cbl', 'language': 'COBOL', 'transcription': '       01 REC.\n          05 VALUE-X PIC X.'}
    sub = {'name':'calc.cbl','language':'COBOL','transcription':'       PROCEDURE DIVISION USING REQUEST RESPONSE.\n           GOBACK.'}
    assert A._layer_role(copy, {}, []) == ('Data', 'Record layout copybook')
    assert A._layer_role(sub, {}, []) == ('Application', 'Called subprogram')


def test_supplied_jcl_and_standard_utilities_are_not_missing_custom_source():
    from core.report.template_docx import _missing
    arts = [{'name':'job.jcl','language':'JCL','transcription':'//S EXEC PGM=IEFBR14\n//T EXEC PGM=IEBGENER'}]
    stack = {}
    _infer_stack(stack, arts + [{'name':'calc.cbl','language':'COBOL','transcription':'SELECT INFILE ASSIGN TO PAYIN.'}],
                 {'sql_server': [], 'db2': False}, {'platforms': []}, EmptyStore())
    assert 'JCL and scheduler not provided' not in str(stack)
    assert 'JCL supplied: job.jcl' in str(stack)
    assert not _missing({'missing':[{'category':'missing_code','name':n} for n in ('IEFBR14','IEBGENER','IEFBR14DCB')]}, arts)


def test_support_prose_does_not_invent_staffing_shortages():
    from core.report import plain as P
    assert 'unconfirmed' in P.what('SUP-SKILLS', 'RPG specialist skills required')
    assert 'few people' not in P.what('SUP-SKILLS')


def test_jcl_continuation_does_not_invent_program_name():
    from core.cobol.parser import parse_jcl
    parsed = parse_jcl('//S EXEC PGM=IEFBR14\n// DCB=(RECFM=FB,LRECL=80)', 'job.jcl')
    assert any(r['target'] == 'program:IEFBR14' for r in parsed['relations'])
    assert 'IEFBR14DCB' not in str(parsed)
