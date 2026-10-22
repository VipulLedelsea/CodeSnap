from types import SimpleNamespace

from core import deepdive as DD
from core import exec_summary as ES
from test_deepdive import art, store  # noqa: F401  (fixtures)

INDEX = {"A.cbl:F1": {"lines": [3, 3], "statement": "x"}, "A.cbl:F2": {"lines": [4, 5], "statement": "y"}}


def test_citations_that_match_no_finding_are_removed():
    dropped = []
    out = ES.check_citations("Pays twice [A.cbl:F1]. Odd [A.cbl:F9] and [B.cbl:F1, F2].", INDEX, dropped)
    assert "[A.cbl:F1]" in out and "F9" not in out and "B.cbl" not in out
    assert sorted(dropped) == ["A.cbl:F9", "B.cbl:F1", "B.cbl:F2"]


def test_style_has_no_dashes_icons_or_markdown():
    out = ES.clean_style("Risk — high – it spans 0–12 months ✅ **bold** `x`")
    assert not any(c in out for c in "–—✅*`")
    assert "0 to 12 months" in out


def test_a_risk_with_no_real_finding_is_not_reported():
    data = {"verdict": "Replace", "provisional": False, "what_it_does": "Pays aid [A.cbl:F1].", "reasons": ["a", "b", "c", "d"],
            "risks": [{"text": "Real", "severity": "High", "facts": ["A.cbl:F2"]},
                      {"text": "Invented", "severity": "High", "facts": ["Z.cbl:F1"]},
                      {"text": "Odd rating", "severity": "Huge", "facts": ["A.cbl:F1"]}],
            "lifecycle": "Not verified", "lifecycle_verified": False, "decision": "d", "confidence": "high",
            "confidence_and_coverage": "c"}
    out = ES.validate(data, INDEX)
    assert [r["text"] for r in out["risks"]] == ["Real", "Odd rating"]
    assert out["risks"][1]["severity"] == "Medium" and len(out["reasons"]) == 3
    assert out["dropped_citations"] == ["Z.cbl:F1"]


def test_an_unknown_verdict_makes_it_provisional():
    out = ES.validate({"verdict": "Burn it", "risks": []}, INDEX)
    assert out["verdict"] is None and out["provisional"]


def test_paragraphs_follow_the_prompt_order_and_cite_facts():
    s = ES.validate({"verdict": "Remediate", "provisional": True, "what_it_does": "Calculates fees.", "reasons": ["a", "b", "c"],
                     "risks": [{"text": "Fee overflow", "severity": "High", "facts": ["A.cbl:F1"]}], "lifecycle": "Versions unknown",
                     "lifecycle_verified": False, "decision": "Approve a fix", "confidence": "medium",
                     "confidence_and_coverage": "Medium confidence.", "needed_to_finalize": ["volumes"]}, INDEX)
    p = ES.paragraphs(s)
    assert p[0].startswith("Verdict: Remediate (provisional).") and "volumes" in p[0]
    assert any(x.startswith("Risk 1, High: Fee overflow [A.cbl:F1]") for x in p)
    assert any("not verified" in x for x in p) and p[-1].startswith("Confidence medium")


def _review(store, a, facts):
    store.set_meta("deepdive", {str(a["id"]): {"name": a["name"], "hash": DD._hash(a["transcription"]), "prompt_version": DD.PROMPT_VERSION,
                                               "purpose": "Pays aid.", "file_role": "business_logic", "run_mode": "batch",
                                               "facts": facts, "unknowns": []}})


class Client:
    def __init__(self, data):
        self.data, self.calls, self.messages = data, [], self

    def create(self, **kw):
        self.calls.append(kw)
        return SimpleNamespace(content=[SimpleNamespace(type="tool_use", name="record_exec_summary", input=self.data)],
                               usage=SimpleNamespace(input_tokens=10, output_tokens=5), stop_reason="tool_use")


def test_run_saves_a_checked_summary_that_goes_stale_when_findings_change(store):  # noqa: F811
    a = art(store, "AidPaymentController.cs")
    name = a["name"]
    _review(store, a, [{"id": 1, "fact_id": "F1", "category": "defect", "severity": "high", "basis": "observed", "lines": [28, 28],
                        "statement": "Approve sets the batch status."}])
    data = {"verdict": "Remediate", "provisional": True, "what_it_does": f"Approves batches [{name}:F1] [{name}:F7].",
            "reasons": ["a", "b", "c"], "risks": [{"text": "Unchecked approval", "severity": "High", "facts": [f"{name}:F1"]}],
            "lifecycle": "Not verified", "lifecycle_verified": False, "decision": "Fix it", "confidence": "low",
            "confidence_and_coverage": "One file."}
    client = Client(data)
    out = ES.run(store, client, model="claude-sonnet-5")
    call = client.calls[0]
    assert call["system"] == ES.SYSTEM and "Write only output item 1" in call["messages"][0]["content"]
    assert f"[{name}:F1]" in call["messages"][0]["content"]
    assert f"[{name}:F7]" not in out["what_it_does"] and out["dropped_citations"] == [f"{name}:F7"]
    assert ES.current(store)["verdict"] == "Remediate"
    _review(store, a, [{"id": 1, "fact_id": "F1", "category": "defect", "severity": "low", "basis": "observed", "lines": [28, 28],
                        "statement": "Changed after the summary was written."}])
    assert ES.current(store) is None


def test_no_summary_without_reviewed_files(store):  # noqa: F811
    assert ES.run(store, Client({}), model="m") is None


def test_the_review_keeps_the_new_per_file_fields():
    props = DD.ANALYSIS_TOOL["input_schema"]["properties"]
    for k in ("file_role", "run_mode", *DD.ROLLUPS):
        assert k in props
    assert "fact_id" in props["facts"]["items"]["properties"]
    assert DD._enum("nonsense", DD.RUN_MODES) == "not shown" and DD._enum("batch", DD.RUN_MODES) == "batch"


def test_the_summary_is_its_own_file_and_the_report_only_points_to_it(store):  # noqa: F811
    import io
    import docx
    from core.report import docx_bytes, summary_doc
    a = art(store, "AidPaymentController.cs")
    name = a["name"]
    _review(store, a, [{"id": 1, "fact_id": "F1", "category": "defect", "severity": "high", "basis": "observed", "lines": [28, 28],
                        "statement": "Approve sets the batch status."}])
    ES.run(store, Client({"verdict": "Remediate", "provisional": True, "what_it_does": "",
                          "reasons": ["a", "b", "c"], "risks": [{"text": "Unchecked approval", "severity": "High", "facts": [f"{name}:F1"]}],
                          "lifecycle": "Not verified", "lifecycle_verified": False, "decision": "Fix it", "confidence": "low",
                          "confidence_and_coverage": "One file."}), model="m")
    report = docx.Document(io.BytesIO(docx_bytes(store, rescan=True)))
    text = "\n".join(p.text for p in report.paragraphs)
    assert f"{store.info['slug']}_executive_summary.docx" in text                       # the report points to the separate file
    assert "Verdict: Remediate (" not in text and "Unchecked approval" not in text     # and no longer carries the summary itself
    assert "Evidence for the key findings" not in text                                  # the evidence appendix moved with the summary
    own = docx.Document(io.BytesIO(summary_doc.docx_bytes(store)))
    body = "\n".join(p.text for p in own.paragraphs)
    cells = "\n".join(c.text for t in own.tables for r in t.rows for c in r.cells)
    assert "Remediate (provisional)" in body and "Unchecked approval" in cells
    assert "<insert business purpose and users information here>" in body              # an unknown is asked for, not guessed
    assert "Approve sets the batch status." in cells and f"[{name}:F1]" not in body + cells   # tags become plain references


def test_text_is_cut_to_whole_sentences_within_the_limit_and_unknowns_become_inserts():
    assert ES.brief("One two three. Four five six seven eight nine. Ten eleven.", 8) == "One two three."
    assert ES.brief("A short one that fits fine.", 10) == "A short one that fits fine."    # within 1.4 times the limit it is kept whole
    out = ES.validate({"verdict": "Replace", "what_it_does": "", "reasons": ["a"], "risks": [], "lifecycle": "<Insert hosting cost here>",
                       "decision": " ", "confidence_and_coverage": "x", "needed_to_finalize": ["<insert user volume here>"]}, {})
    assert out["what_it_does"] == "<insert business purpose and users information here>"
    assert out["lifecycle"] == "<insert hosting cost information here>" and out["reasons"][1].startswith("<insert reason")
    assert out["needed_to_finalize"] == ["<insert user volume information here>"]
    assert "under 300 words" in ES.SCOPE and "<insert {specific item} information here>" in ES.SCOPE


def test_a_bracket_listing_several_files_is_checked_fact_by_fact():
    dropped = []
    out = ES.check_citations("Rules clash [A.cbl:F1, B.cbl:F2, B.cbl:F9] and more [A.cbl:F3, F4].", {**INDEX, "B.cbl:F2": {}}, dropped)
    assert out == "Rules clash [A.cbl:F1] [B.cbl:F2] and more ." or out.startswith("Rules clash [A.cbl:F1] [B.cbl:F2] and more")
    assert dropped == ["B.cbl:F9", "A.cbl:F3", "A.cbl:F4"]
    assert ES.cited_ids("x [A.cbl:F1] [B.cbl:F2]") == ["A.cbl:F1", "B.cbl:F2"]


def test_a_long_sentence_is_trimmed_but_keeps_its_citations():
    long = "Nine programs each hardcode the same district rules, so the same payment gets different outcomes, depending on which program touches it first and which district it comes from [A.cbl:F1]."
    out = ES.brief(long, 12)
    assert out.endswith("[A.cbl:F1]") and len(out.replace("[A.cbl:F1]", "").split()) <= 14


def test_the_consistency_read_ignores_items_it_calls_consistent():
    from core import report_consistency as RC
    text = "Quote number one is here. Quote number two is here."
    items = [{"first_quote": "Quote number one is here.", "second_quote": "Quote number two is here.", "why": "Actually consistent - skip"},
             {"first_quote": "Quote number one is here.", "second_quote": "Quote number two is here.", "why": "They give different totals."}]
    assert len(RC.verified(items, text)) == 1


class Verdicts(Client):
    """Gives the listed verdicts one call at a time."""

    def __init__(self, words, name):
        super().__init__(None)
        self.words, self.name = list(words), name

    def create(self, **kw):
        self.data = {"verdict": self.words.pop(0) if len(self.words) > 1 else self.words[0], "provisional": False, "what_it_does": "Pays.",
                     "reasons": ["a", "b", "c"], "risks": [{"text": "r", "severity": "High", "facts": [f"{self.name}:F1"]}],
                     "lifecycle": "l", "lifecycle_verified": False, "decision": "d", "confidence": "low", "confidence_and_coverage": "c"}
        return super().create(**kw)


def _setup(store, code):  # noqa: F811
    a = art(store, "AidPaymentController.cs")
    _review(store, a, [{"id": 1, "fact_id": "F1", "category": "defect", "severity": "high", "basis": "observed", "lines": [28, 28],
                        "statement": "Approve skips a check."}])
    store.set_meta("assessment", {"verdict": {"code": code, "label": code.title(), "meaning": "m"}})
    return a["name"]


def test_the_summary_verdict_has_to_match_the_reports_own_disposition(store):  # noqa: F811
    name = _setup(store, "retain")
    assert ES.report_verdict(store)[0] == "Remediate"          # keep and fix, with a high finding open
    c = Verdicts(["Replace", "Remediate"], name)
    out = ES.run(store, c, model="m")
    assert out["verdict"] == "Remediate" and len(c.calls) == 2 and "model_verdict" not in out
    assert "recommend Remediate" in c.calls[0]["messages"][0]["content"]


def test_a_verdict_that_stays_different_is_replaced_and_the_difference_is_kept(store):  # noqa: F811
    name = _setup(store, "rearchitect")
    out = ES.run(store, Verdicts(["Retire"], name), model="m")
    assert out["verdict"] == "Replace" and out["model_verdict"] == "Retire"


def test_a_correction_becomes_a_short_note_without_internal_finding_numbers():
    from core.report_review import note_for
    assert note_for("Finding 150 confirms MUNPAY.cbl uses PAYIN/PAYOUT; findings 252/254 confirm MUNREV does too") == ""
    assert note_for("Findings confirm PL/I, RPG and COBOL") == ""
    assert note_for("Finding 1 shows 32 error pages (H00-H31), not 50") == "The code shows 32 error pages (H00-H31), not 50."
    assert "Finding" not in note_for("The rate is 4 percent, not 5 percent (Finding 12).")


def test_a_bare_insert_is_named_by_its_column_and_row():
    from docx import Document
    from core.report.template_docx import _label_bare_inserts
    d = Document()
    t = d.add_table(rows=2, cols=3)
    for i, h in enumerate(["Data entity", "Retention requirement", "Data classification"]):
        t.cell(0, i).text = h
    t.cell(1, 0).text = "MUN_ACCOUNT0"
    t.cell(1, 1).text = "<insert information here>"
    t.cell(1, 2).text = "Public"
    _label_bare_inserts(SimpleNamespace(d=d))
    assert "<insert" in t.cell(1, 1).text and "information here>" in t.cell(1, 1).text
    assert "<insert information here>" not in t.cell(1, 1).text and "retention requirement" in t.cell(1, 1).text.lower()
    assert t.cell(1, 2).text == "Public"


def test_the_summary_marks_statements_beside_a_gap_and_does_not_claim_there_were_no_reading_concerns():
    from core.report.summary_doc import _confidence
    level, text = _confidence("high", "All 12 submitted files were reviewed with no reading concerns, giving high coverage.",
                              {"07_MUNREV.cbl": [(919, 925)], "12_MUNINQ.bms": [(475, 481)]})
    assert level == "medium" and "no reading concerns" not in text and text.startswith("The copy of 2 files has gaps")
    assert _confidence("high", "All good.", {}) == ("high", "All good.")


def test_evidence_rows_read_cleanly_in_the_executive_summary():
    from core import exec_summary as ES
    index = {"a.cbl:F1": {"lines": [7, 7], "severity": "high", "category": "high",
                          "statement": "Four IFs are not exclusive, so if several are true (e.g. WS-RESULT is moved four times in a row) the last wins."},
             "a.cbl:F2": {"lines": [10, 20], "severity": "medium", "category": "data_integrity", "statement": "A short finding."}}
    s = {"verdict": "Retain", "risks": [{"severity": "High", "text": "x", "facts": ["a.cbl:F1", "a.cbl:F2"]}]}
    rows = ES.evidence_rows(s, index)
    assert rows[0][1] == "7" and rows[0][2] == "Finding" and rows[1][1] == "10-20" and rows[1][2] == "Data integrity"
    assert rows[0][3].count("(") == rows[0][3].count(")") and "e.g.." not in rows[0][3]
    short = ES.brief("Word " * 80 + "(e.g. something here and more that runs on", 10)
    assert short.count("(") == short.count(")") and short.endswith(("…", "."))
