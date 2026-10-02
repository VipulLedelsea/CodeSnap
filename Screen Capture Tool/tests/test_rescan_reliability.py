"""Offline regressions for transcription completeness, report credibility and timely UI updates."""
import io
import json
import threading
import subprocess
from pathlib import Path
from types import SimpleNamespace

import pytest
from core.model import ProgramStore
from core import analysis


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setenv('CODESNAP_PROGRAMS', str(tmp_path/'programs'))
    with ProgramStore.create('Rescan fixture') as st:
        yield st


def function(name):
    source = Path('src/webapp/static/app.js').read_text()
    start = source.index('async function '+name+'(')
    return source[start:source.index('\n}', start)+2]


def node(script):
    subprocess.run(['node', '-e', script], check=True, capture_output=True, text=True)


def test_incomplete_screenshot_stops_finalization(tmp_path, monkeypatch):
    a,b=tmp_path/'a.png',tmp_path/'b.png'
    a.write_bytes(b'a');b.write_bytes(b'b')
    calls=[]
    def extract(client,path):
        if path==b: raise RuntimeError('failed screenshot')
        return {'raw':'say 1','verify':{}}
    monkeypatch.setattr(analysis,'extract_structured',extract)
    monkeypatch.setattr(analysis,'synthesize_final',lambda *a: calls.append('final'))
    with pytest.raises(RuntimeError,match='failed screenshot'):
        analysis.analyse_incremental(None,[a,b],tmp_path/'.cache')
    assert not calls
    assert analysis.cache_path_for(a,tmp_path/'.cache').exists()
    assert not analysis.cache_path_for(b,tmp_path/'.cache').exists()


def test_transcription_truncation_is_not_cached(tmp_path, monkeypatch):
    path=tmp_path/'a.png';path.write_bytes(b'fake')
    client=SimpleNamespace(messages=SimpleNamespace(create=lambda **kw: SimpleNamespace(
        stop_reason='max_tokens',content=[SimpleNamespace(text='{"raw_transcription":["say 1"]}')])) )
    with pytest.raises(ValueError,match='truncated'):
        analysis.extract_to_cache(client,path,tmp_path/'.cache')
    assert not analysis.cache_path_for(path,tmp_path/'.cache').exists()


def test_fallback_and_merge_preserve_observed_indentation():
    raw='      say 1\n      say 2\n'
    assert analysis._normalize_extract(raw)['raw'].startswith('      say 1')
    code,*_=analysis.merge_verified([raw])
    assert code.startswith('      say 1')


def test_unmeasured_sources_block_release(store):
    from core.report import quality
    from core import deepdive
    aid=store.add_artifact('policy.py',language='Python',transcription='print(1)')
    store.set_meta('deepdive',{str(aid):{'hash':deepdive._hash('print(1)'),'facts':[]}})
    assert any('source completeness remains unconfirmed' in b for b in quality.summary(store)['blockers'])


def test_rebuild_failure_rolls_back_source_and_analysis(store,tmp_path,monkeypatch):
    from PIL import Image
    from core.model import restitch
    shot=tmp_path/'a.png';Image.new('RGB',(10,10),'white').save(shot)
    eid=store.add_evidence(shot)
    aid=store.add_artifact('policy.py',language='Python',transcription='def old():\n    return 1',evidence_ids=[eid])
    store.upsert_entity('function','old',artifact_id=aid)
    store.text_cache_dir.joinpath(store.evidence(eid)['sha256']+'.md').write_text('def new():\n    return 2')
    monkeypatch.setattr(restitch,'apply_structure',lambda *a,**kw: (_ for _ in ()).throw(RuntimeError('save failed')))
    with pytest.raises(RuntimeError,match='save failed'):
        restitch.rebuild_saved_capture(store,aid)
    assert store.current_artifact('policy.py')['id']==aid
    assert store.entity_by_key('function:old')
    assert len(store.artifacts(current_only=False))==1


def test_report_privacy_settings_are_request_local(store,monkeypatch):
    from core.report import content, settings
    from core.assess import run_assessment
    store.add_artifact('policy.py',language='Python',transcription='print(1)')
    run_assessment(store)
    settings.save(store,{'privacy_obligations':'FIRST POLICY'})
    first=content.build(store,rescan=False)
    settings.save(store,{'privacy_obligations':'SECOND POLICY'})
    second=content.build(store,rescan=False)
    assert 'FIRST POLICY' not in json.dumps(second)
    assert 'SECOND POLICY' not in json.dumps(first)
    assert not hasattr(content,'_PRIVACY')


def test_word_rating_mode_is_thread_local():
    from core.report import rationale
    gate=threading.Barrier(2)
    observed=[]
    def words():
        with rationale.word_ratings():
            gate.wait();observed.append(rationale._word_mode.get());gate.wait()
    thread=threading.Thread(target=words);thread.start();gate.wait()
    assert not rationale._word_mode.get()
    gate.wait();thread.join(2)
    assert observed==[True] and not thread.is_alive()


def test_html_and_word_share_issued_text_and_disposition(store):
    import docx
    from core.report import package
    from core.model.ingest import ingest_artifact
    from html.parser import HTMLParser
    class TextParser(HTMLParser):
        def __init__(self): super().__init__();self.text=[]
        def handle_data(self,data): self.text.append(data)
    aid=store.add_artifact('policy.py',language='Python',transcription='def calculate():\n    return 1')
    ingest_artifact(store,None,aid)
    result=package(store,pdf=False)
    output=Path(result['dir']);slug=store.info['slug']
    word=docx.Document(output/f'{slug}_assessment_report.docx')
    parser=TextParser();parser.feed((output/f'{slug}_assessment_report.html').read_text())
    html=' '.join(parser.text)
    norm=lambda text:' '.join(text.split())
    text=norm(html)
    for value in [p.text for p in word.paragraphs]+[c.text for t in word.tables for row in t.rows for c in row.cells]:
        if value.strip(): assert norm(value) in text
    assert 'Keep as is; routine maintenance only.' not in text
    assert 'Draft for review' in text and 'Blocks issue:' in text
    assert result['verdict']['label'] in text


def test_poll_failure_schedules_retry_and_latest_response_wins():
    source=Path('src/webapp/static/app.js').read_text()
    signature='function _programRenderSignature('+source.split('function _programRenderSignature(',1)[1].split('\nasync function loadProgram',1)[0]
    node(signature+function('loadProgram')+r'''
const assert=require('node:assert/strict');
let _program='A',_progSig='';const _localRebuilds=new Map(),elements={};
const $=id=>elements[id]||={style:{},textContent:'',innerHTML:''};
const document={querySelectorAll:()=>[]};
const renderProgress=()=>{},renderCaptureBar=()=>{},toast=()=>{};
let retries=0;const scheduleProgramPoll=()=>retries++;
const _fileRow=row=>row.name;
let _json=async()=>{throw Error('temporary refresh failure')};
const data=name=>({program:{},progress:{total:1},artifacts:[{id:1,name,state:'done'}],coverage:{missing:[]}});
(async()=>{
 await loadProgram(true);assert.equal(retries,1);
 const requests=[];_json=()=>new Promise(resolve=>requests.push(resolve));
 const old=loadProgram(true),recent=loadProgram();
 requests[1](data('current'));await recent;
 requests[0](data('stale'));await old;
 assert.equal($('progFiles').innerHTML,'current');
})().catch(e=>{console.error(e);process.exitCode=1});
''')


def test_secondary_panel_rejects_late_program_response():
    node(function('loadAssessment')+r'''
const assert=require('node:assert/strict');let _program='A',_assess=null,_assessInputs={};
const requests=[],rendered=[];
const _json=()=>new Promise(resolve=>requests.push(resolve));
const renderAssessment=()=>rendered.push(_assess.name),$=()=>({});
(async()=>{
 const old=loadAssessment();_program='B';const current=loadAssessment();
 requests[1]({assessment:{name:'B'},inputs:{}});await current;
 requests[0]({assessment:{name:'A'},inputs:{}});await old;
 assert.deepEqual(rendered,['B']);assert.equal(_assess.name,'B');
})().catch(e=>{console.error(e);process.exitCode=1});
''')


def test_report_settings_save_stays_bound_to_original_program():
    node(function('loadReportSettings')+r'''
const assert=require('node:assert/strict');let _program='A',click,requested;
const box={innerHTML:'',querySelectorAll:()=>[{dataset:{k:'client'},value:'CLIENT A'}]};
const $=id=>id==='progReportForm'?box:{addEventListener:(event,fn)=>click=fn};
const escapeHtml=x=>x,_attr=x=>x,toast=()=>{};
const _json=async(url,options)=>{if(options)requested=url;return {fields:[],values:{}}};
(async()=>{
 await loadReportSettings();_program='B';await click();
 assert.equal(requested,'/api/programs/A/report-settings');
})().catch(e=>{console.error(e);process.exitCode=1});
''')


def test_literal_spacing_survives_escaped_and_unclosed_quotes():
    from core.text import normalized_line
    assert normalized_line(r'print("A\"  B")') != normalized_line(r'print("A\" B")')
    assert normalized_line("say 'A  B [CUT OFF]") != normalized_line("say 'A B [CUT OFF]")
    assert normalized_line("say  'A''  B'") == normalized_line("say 'A''  B'")


def test_busy_poll_interval_and_resume_refresh():
    source=Path('src/webapp/static/app.js').read_text()
    schedule='function scheduleProgramPoll('+source.split('function scheduleProgramPoll(',1)[1].split('\nlet _display',1)[0]
    resume='document.addEventListener("visibilitychange",'+source.split('document.addEventListener("visibilitychange",',1)[1]
    node(r'''
const assert=require('node:assert/strict');let _progPoll=null,_view='review',_program='A',delay,listener,refreshes=0;
const clearTimeout=()=>{},setTimeout=(fn,ms)=>{delay=ms;return 1};
const document={hidden:false,addEventListener:(event,fn)=>listener=fn};
const loadProgram=poll=>{assert.equal(poll,true);refreshes++};
'''+schedule+resume+r'''
scheduleProgramPoll({progress:{to_go:1}});assert.equal(delay,3000);
scheduleProgramPoll({progress:{to_go:0},session_running:false});assert.equal(delay,10000);
listener();assert.equal(refreshes,1);
document.hidden=true;listener();assert.equal(refreshes,1);
''')


def test_model_revision_refreshes_open_assessment_after_analysis():
    source=Path('src/webapp/static/app.js').read_text()
    signature='function _programRenderSignature('+source.split('function _programRenderSignature(',1)[1].split('\nasync function loadProgram',1)[0]
    node(signature+function('loadProgram')+r'''
const assert=require('node:assert/strict');let _program='A',_progSig='',revision='1',refreshes=0;
const _localRebuilds=new Map(),elements={};
const $=id=>elements[id]||={style:{},textContent:'',innerHTML:'',open:id==='progAssessBox'};
const document={querySelectorAll:()=>[]};
const renderProgress=()=>{},renderCaptureBar=()=>{},scheduleProgramPoll=()=>{},loadAssessment=()=>refreshes++;
const _fileRow=row=>row.name;
const _json=async()=>({program:{},revision,progress:{total:1},artifacts:[{id:1,name:'same file',state:'done'}],coverage:{missing:[]}});
(async()=>{
 await loadProgram(true);assert.equal(refreshes,1);
 revision='2';await loadProgram(true);assert.equal(refreshes,2);
 await loadProgram(true);assert.equal(refreshes,2);
})().catch(e=>{console.error(e);process.exitCode=1});
''')


def test_source_word_export_preserves_literal_lines():
    source = '    first\n# compiler directive\n- literal source\n  last  '
    doc = analysis.build_docx({'extracted_text': source})
    assert '\n'.join(p.text for p in doc.paragraphs) == source
    assert all(p.style.name == 'Normal' for p in doc.paragraphs)


def test_unclosed_quote_preserves_trailing_literal_spaces():
    from core.text import normalized_line
    assert normalized_line("  say 'A  ") == "say 'A  "
    assert normalized_line("say 'A  ") != normalized_line("say 'A ")
    assert normalized_line("  say 'A'  ") == "say 'A'"


def test_empty_cached_frame_cannot_publish_partial_rebuild(store,tmp_path):
    from PIL import Image
    from core.model.restitch import rebuild_saved_capture
    ids=[]
    for index,text in enumerate(('say 1','   \n')):
        shot=tmp_path/f'{index}.png';Image.new('RGB',(10+index,10),'white').save(shot)
        eid=store.add_evidence(shot);ids.append(eid)
        (store.text_cache_dir/(store.evidence(eid)['sha256']+'.md')).write_text(text)
    aid=store.add_artifact('routine.rexx',language='REXX',transcription='say 1\nsay 2',evidence_ids=ids)
    with pytest.raises(ValueError,match='Screenshot 2 has no cached readable text'):
        rebuild_saved_capture(store,aid)
    assert store.current_artifact('routine.rexx')['id']==aid
    assert len(store.artifacts(current_only=False))==1


def test_old_comparison_results_require_refresh(store,tmp_path):
    from core.report import quality
    store.add_artifact('routine.rexx',language='REXX',transcription="say 'A  '")
    original=tmp_path/'original.rexx';original.write_text("say 'A  '")
    benchmark=quality.measure(store,{'routine.rexx':original})
    benchmark.pop('comparison_revision')
    store.set_meta('transcription_benchmark',benchmark)
    assert quality.current_benchmark(store)['files']==[]
    assert quality.current_benchmark(store)['stale']==['routine.rexx']
    assert any('benchmark needs refresh' in b for b in quality.summary(store)['blockers'])
    quality.measure(store,{'routine.rexx':original})
    assert not quality.current_benchmark(store)['stale']
    assert quality.current_benchmark(store)['files'][0]['matched']==1


def test_slow_status_poll_is_not_superseded_by_interval():
    node(r'''
const assert=require('node:assert/strict');let release,calls=0,renders=0;
const fetch=()=>{calls++;return new Promise(resolve=>release=resolve)};
const renderStatus=()=>renders++,renderFlow=()=>{},setRunning=()=>{},startAnalyzing=()=>{},stopAnalyzing=()=>{};
let lastEventCount='[]',_program='';const loadReports=()=>{};
'''+function('pollStatus')+r'''
(async()=>{
  const first=pollStatus();await pollStatus();await pollStatus();
  assert.equal(calls,1);assert.equal(renders,0);
  release({json:async()=>({events:[],running:false})});await first;
  assert.equal(renders,1);assert.equal(pollStatus.pending,false);
  const second=pollStatus();assert.equal(calls,2);
  release({json:async()=>{throw Error('offline')}});await second;
  assert.equal(pollStatus.pending,false);
})();
''')


def test_slow_program_refresh_survives_repeated_background_polls():
    node(function('loadProgram')+r'''
const assert=require('node:assert/strict');let _program='A',release,calls=0,retries=0;
const $=()=>({style:{}}),toast=()=>{},scheduleProgramPoll=()=>retries++;
const _json=()=>{calls++;return new Promise((resolve,reject)=>release=reject)};
(async()=>{
 const first=loadProgram(true);await loadProgram(true);await loadProgram(true);
 assert.equal(calls,1);release(Error('offline'));await first;
 assert.equal(retries,1);assert.equal(loadProgram.pending,null);
 const next=loadProgram(true);assert.equal(calls,2);release(Error('offline'));await next;
 assert.equal(retries,2);
})().catch(e=>{console.error(e);process.exitCode=1});
''')
