import io
from types import SimpleNamespace

from docx import Document

from core import deepdive as D
from core import report_consistency as RC
from core import report_review as RR

CODE = "MOVE A TO B\nCALL 'CUSTMAST' USING REC\nSTOP RUN"


def make(paras, table=None):
    d = Document()
    d.add_heading("1. Summary", 1)
    for p in paras:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def test_correction_needs_a_quote_from_the_file():
    assert D.correction_backed({"corrected_quote": "CALL 'CUSTMAST' USING REC"}, CODE)
    assert not D.correction_backed({"corrected_quote": "CALL 'OTHER' USING REC"}, CODE)
    assert not D.correction_backed({}, CODE)


def test_prompts_carry_the_new_rules():
    assert "you see one file at a time only" in D.DEEP_SYSTEM
    assert "Comments, headers, names" in D.DEEP_SYSTEM
    assert "misread test" in D.DEEP_SYSTEM
    assert "corrected_quote" in D.REVIEW_SYSTEM
    assert "at least two files" in D.SYNTH_SYSTEM
    assert D.DEFAULT_REVIEW_MODEL != "claude-haiku-4-5"


def test_unsupported_needs_a_contradicting_quote_or_a_verified_absence():
    claim = "The program writes every record to INVMAST."
    assert RR.gate({"verdict": "unsupported"}, claim, text=CODE) == "cannot_check"
    assert RR.gate({"verdict": "unsupported", "quote": "STOP RUN"}, claim, text=CODE) == "unsupported"
    assert RR.gate({"verdict": "unsupported", "absent_term": "INVMAST"}, claim, text=CODE) == "unsupported"
    assert RR.gate({"verdict": "unsupported", "absent_term": "CUSTMAST"}, claim, text=CODE) == "cannot_check"
    assert RR.gate({"verdict": "unsupported", "absent_term": "ZZZZ9999"}, claim, text=CODE) == "cannot_check"


def test_supported_needs_a_quote():
    assert RR.gate({"verdict": "supported"}, "x", text=CODE) == "cannot_check"
    assert RR.gate({"verdict": "supported", "quote": "MOVE A TO B"}, "x", text=CODE) == "supported"


def test_cross_file_verdicts_need_a_real_finding_number():
    assert RR.gate({"verdict": "unsupported"}, "x", findings=5) == "cannot_check"
    assert RR.gate({"verdict": "unsupported", "finding": 9}, "x", findings=5) == "cannot_check"
    assert RR.gate({"verdict": "unsupported", "finding": 3}, "x", findings=5) == "unsupported"
    assert RR.gate({"verdict": "bogus", "finding": 3}, "x", findings=5) == "cannot_check"


def test_claim_naming_two_files_is_not_checked_against_one():
    raw = make(["PAYROLL.cbl passes the gross amount to TAXCALC.cbl before every record is written out."])
    cl = RR.extract_claims(raw, ["PAYROLL.cbl", "TAXCALC.cbl"])
    assert sorted(cl[0]["files"]) == ["PAYROLL.cbl", "TAXCALC.cbl"]
    assert len(cl[0]["files"]) != 1


def test_claim_carries_its_paragraph_for_context():
    raw = make(["PAYROLL.cbl opens the customer master file at the start of each run. It never closes the file after the last record is written."])
    cl = RR.extract_claims(raw, ["PAYROLL.cbl"])
    assert all("PAYROLL.cbl opens the customer" in c["para"] for c in cl)
    assert "(paragraph: " in RR.claim_line(2, cl[1])


def test_summary_does_not_claim_everything_was_checked():
    store = SimpleNamespace(get_meta=lambda k: {"report_review": {"verdicts": {
        "a": {"verdict": "supported"}, "b": {"verdict": "cannot_check"}, "c": {"verdict": "unsupported"}}}}.get(k))
    line = RR.summary_line(store)
    assert "Every statement" not in line and line.startswith("3 statements") and "could not be checked" in line


def test_consistency_items_must_quote_the_report():
    text = "Effort is 5 to 14 person-weeks.\nLater: effort is 34 to 72 person-weeks."
    items = [{"first_quote": "Effort is 5 to 14 person-weeks.", "second_quote": "effort is 34 to 72 person-weeks.", "topic": "effort", "kind": "range", "why": "w"},
             {"first_quote": "Effort is 5 to 14 person-weeks.", "second_quote": "made up sentence not in report", "topic": "t", "kind": "number", "why": "w"},
             {"first_quote": "effort is 34 to 72 person-weeks.", "second_quote": "Effort is 5 to 14 person-weeks.", "topic": "effort", "kind": "range", "why": "dup"},
             {"first_quote": "Effort is 5 to 14 person-weeks.", "second_quote": "Effort is 5 to 14 person-weeks.", "topic": "same", "kind": "number", "why": "w"}]
    out = RC.verified(items, text)
    assert len(out) == 1 and out[0]["topic"] == "effort"


def test_report_text_has_headings_and_table_rows():
    raw = make(["Body sentence about the program."], table=[["Risk ID", "Title"], ["R1", "Hard-coded password"]])
    text = RC.report_text(raw)
    assert "## 1. Summary" in text and "R1 | Hard-coded password" in text


def test_consistency_run_stores_only_verified_items(monkeypatch):
    raw = make(["Effort is 5 to 14 person-weeks.", "Elsewhere the effort is 34 to 72 person-weeks."])
    block = SimpleNamespace(type="tool_use", name="record_inconsistencies", input={"items": [
        {"topic": "effort", "kind": "range", "first_quote": "Effort is 5 to 14 person-weeks.",
         "second_quote": "the effort is 34 to 72 person-weeks.", "why": "two ranges"},
        {"topic": "x", "kind": "count", "first_quote": "invented one", "second_quote": "invented two", "why": "n"}]})
    msg = SimpleNamespace(content=[block], stop_reason="tool_use")
    monkeypatch.setattr(D, "_call", lambda *a, **k: (msg, 1))
    monkeypatch.setattr(D, "_log", lambda *a, **k: None)
    from core import report_numbers
    monkeypatch.setattr(report_numbers, "truth", lambda st: {})
    monkeypatch.setattr(report_numbers, "check", lambda text, truth: [])  # this test is about quote checking only
    meta = {}
    store = SimpleNamespace(set_meta=lambda k, v: meta.__setitem__(k, v), get_meta=lambda k: meta.get(k))
    res = RC.run(store, object(), raw, model="m")
    assert len(res["items"]) == 1 and RC.items(store)[0]["topic"] == "effort"


def test_placeholder_names_the_missing_information():
    from core.report import template_docx as T
    assert T.placeholder("Client organization") == "<insert client organization information here>"
    assert T.placeholder("RTO") == "<insert RTO information here>"
    assert T.placeholder("Availability, RTO, RPO, batch window and peak volumes (6.3)") == "<insert availability, RTO, RPO, batch window and peak volumes information here>"
    assert T.placeholder("") == "<insert missing information here>"


def test_unknown_cells_are_labelled_by_row_and_column():
    from core.report import template_docx as T
    d = Document()
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["Item", "Current", "Target"], ["Batch window", "Unknown", "Unknown"], ["Peak volumes", "5/s", "Unknown (only the code was read)"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    kv = d.add_table(rows=2, cols=2)
    for r, row in enumerate([["Field", "Value"], ["Business owner", "Unknown"]]):
        for c, v in enumerate(row):
            kv.cell(r, c).text = v
    T._label_unknowns(SimpleNamespace(d=d))
    assert t.cell(1, 1).text == "<insert batch window current information here>"
    assert t.cell(1, 2).text == "<insert batch window target information here>"
    assert t.cell(2, 1).text == "5/s"
    assert t.cell(2, 2).text == "<insert peak volumes target information here> (only the code was read)"
    assert kv.cell(1, 1).text == "<insert business owner information here>"
