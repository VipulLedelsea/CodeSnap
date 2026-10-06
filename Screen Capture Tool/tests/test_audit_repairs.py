"""Regression coverage for the October 1 audit repairs; providers and programs are synthetic."""
import os
import sqlite3
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from core.model import ProgramStore
from core.model.ingest import complete_capture


@pytest.fixture
def store(tmp_path, monkeypatch):
    monkeypatch.setenv('CODESNAP_PROGRAMS', str(tmp_path / 'programs'))
    with ProgramStore.create('Audit fixture') as st:
        yield st


def image(tmp_path, name='one.png'):
    p = tmp_path / name
    Image.new('RGB', (20, 20), 'red').save(p)
    return p


def test_rebuild_crashes_with_running_worker(store, monkeypatch):
    from webapp import server
    aid = store.add_artifact('a.rexx', transcription='say 1')
    monkeypatch.setattr(server._session, 'running', lambda: True)
    result = server.api_program_artifact_rebuild(store.info['slug'], aid)
    assert result.status_code == 409


def test_pending_replacement_loses_target_after_rename(store, tmp_path):
    old = store.add_artifact('old.rexx', transcription='say 1')
    pending = store.add_pending_capture([image(tmp_path)], recapture_of='old.rexx')
    store.rename_artifact(old, 'renamed.rexx')
    result = complete_capture(store, None, pending, {'code': 'say 2', 'extension': 'rexx', 'code_name': 'replacement'})
    assert len(store.artifacts()) == 1
    assert store.current_artifact('renamed.rexx')['transcription'] == 'say 2'
    assert store.current_artifact('renamed.rexx')['version'] == 2


def test_correction_history_does_not_follow_file_rename(store):
    from core.model import corrections as C
    aid = store.add_artifact('old.rexx', transcription='say 1')
    correction = C.add(store, 'artifact.replace_line', {'artifact': 'old.rexx', 'line': 1,
                                                       'old_text': 'say 1', 'new_text': 'say 2'})
    store.rename_artifact(aid, 'renamed.rexx')
    C.undo(store, correction['id'])
    assert store.current_artifact('renamed.rexx')['transcription'] == 'say 1'


def test_distinct_file_names_overwrite_same_source_path(store):
    one = store.add_artifact('customer order.rexx', transcription='say 1')
    two = store.add_artifact('customer_order.rexx', transcription='say 2')
    a, b = store.artifact(one), store.artifact(two)
    assert a['source_path'] != b['source_path']
    assert (store.sources_dir / a['source_path']).read_text() == 'say 1'
    assert (store.sources_dir / b['source_path']).read_text() == 'say 2'
    assert a['transcription'] == 'say 1'


def test_failed_version_insert_destroys_previous_analysis(store):
    aid = store.add_artifact('original.rexx', transcription='say 1')
    store.upsert_entity('program', 'ORIGINAL', artifact_id=aid)
    assert store.entities()
    with pytest.raises(sqlite3.IntegrityError):
        store.add_artifact('original.rexx', transcription='say 2', evidence_ids=[99999])
    assert store.current_artifact('original.rexx')['id'] == aid
    assert store.entities()


def test_pending_metadata_updates_lose_other_workers_changes(store, tmp_path, monkeypatch):
    a = store.add_pending_capture([image(tmp_path)])
    b = store.add_pending_capture([image(tmp_path, 'two.png')])
    barrier = threading.Barrier(2)
    errors = []
    with ProgramStore.open(store.info['slug']) as first, ProgramStore.open(store.info['slug']) as second:
        def update(st, aid, label):
            try:
                barrier.wait(3)
                st.update_pending(aid, stage=label)
            except Exception as exc:
                errors.append(exc)
        threads = [threading.Thread(target=update, args=(first, a, 'first')),
                   threading.Thread(target=update, args=(second, b, 'second'))]
        for t in threads: t.start()
        for t in threads: t.join(5)
        assert not errors
    pending = store.pending_captures()
    assert sum(bool(v.get('stage')) for v in pending.values()) == 2


def test_retry_releases_an_active_workers_claim(store, tmp_path, monkeypatch):
    from webapp import server
    aid = store.add_pending_capture([image(tmp_path)])
    assert store.claim_pending(aid, f'{os.getpid()}:first-worker')
    monkeypatch.setattr(server, 'api_program_pending_process', lambda slug: {'ok': True})
    response = TestClient(server.app).post(f'/api/programs/{store.info["slug"]}/artifacts/{aid}/retry')
    assert response.status_code == 409
    assert not store.claim_pending(aid, f'{os.getpid()}:second-worker')


def test_live_long_running_claim_is_stolen_after_fifteen_minutes(store, tmp_path):
    aid = store.add_pending_capture([image(tmp_path)])
    assert store.claim_pending(aid, f'{os.getpid()}:first-worker')
    info = store.pending_captures()[aid]
    info['claim']['at'] = time.time() - 901
    store.update_pending(aid, claim=info['claim'])
    assert not store.claim_pending(aid, f'{os.getpid()}:second-worker')


def test_running_worker_is_not_switched_to_new_program(tmp_path, monkeypatch):
    from webapp import session
    monkeypatch.setattr(session, 'PROJECT', tmp_path)
    manager = session.SessionManager()
    manager._proc = SimpleNamespace(pid=os.getpid(), poll=lambda: None)
    spawned = []
    monkeypatch.setattr(session.subprocess, 'Popen', lambda argv, **kw: spawned.append(argv))
    with pytest.raises(RuntimeError, match='switching programs'):
        manager.start(program='different-program')
    assert not spawned
    assert manager.program is None


def test_deep_review_bypasses_exhausted_program_budget(store, monkeypatch):
    from webapp import server
    from core import deepdive
    store.set_meta('budget_usd', 0.01)
    store.log_run('transcribe', model='claude-sonnet-5', input_tokens=1, output_tokens=1, cost=0.02)
    aid = store.add_artifact('a.rexx', transcription='say 1')
    finished = threading.Event()
    observed = []
    fake = SimpleNamespace(messages=SimpleNamespace(create=lambda **kw: observed.append('fake request')))
    monkeypatch.setattr(server, '_client', lambda: fake)
    def run(st, client, **kwargs):
        client.messages.create()
        finished.set()
        return {'errors': []}
    monkeypatch.setattr(deepdive, 'run', run)
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as exc:
        server._deep_start(store.info['slug'], [aid])
    assert exc.value.status_code == 402
    assert not observed


def test_parallel_reads_issue_duplicate_requests_for_one_screenshot(tmp_path, monkeypatch):
    from core import analysis
    shot = image(tmp_path)
    calls = []
    entered, release = threading.Event(), threading.Event()
    def extract(client, path):
        calls.append(path)
        entered.set()
        release.wait(30)            # hold the first read open until the second one has started
        return {'raw': 'say 1', 'corrections': []}
    monkeypatch.setattr(analysis, 'extract_structured', extract)
    errors = []
    def read():
        try: analysis.extract_to_cache(None, shot, tmp_path / '.cache')
        except Exception as exc: errors.append(exc)
    first, second = threading.Thread(target=read), threading.Thread(target=read)
    first.start()
    assert entered.wait(30)
    second.start()
    time.sleep(0.2)                 # let the second read reach the cache / lock before the first one finishes
    release.set()
    for t in (first, second): t.join(30)
    assert not first.is_alive() and not second.is_alive()
    assert not errors and len(calls) == 1


def test_failed_auto_classification_strands_unregistered_capture(tmp_path, monkeypatch):
    import hotkey_capture as hk
    from core import capture, notify, status
    monkeypatch.setattr(status, 'publish', lambda *args, **kwargs: None)
    monkeypatch.setattr(notify, 'notify', lambda *args: None)
    monkeypatch.setattr(capture, 'capture_full_png', lambda: b'fake-frame')
    monkeypatch.setattr(capture, 'clarity_warning', lambda: None)
    class FailedFuture:
        def done(self): return True
        def result(self): raise RuntimeError('fake classification timeout')
    app = hk.App.__new__(hk.App)
    app.running, app.region, app.capture_kind, app.client = True, None, 'auto', None
    app._pool = SimpleNamespace(submit=lambda *args: FailedFuture())
    app._own_window_in_front = lambda: False
    registered = []
    app._register_capture = lambda directory: registered.append(directory)
    app.project_mode, app.idle_stop = False, 0
    app._analyse_burst = lambda *a, **kw: None
    monkeypatch.setattr(hk, 'BURST_MAX_FRAMES', 1)
    app._burst_loop(tmp_path)
    assert not app.running and registered
    assert (tmp_path / '001.png').exists()


def test_report_can_be_final_while_open_items_block_final_issue(store, tmp_path):
    import io
    import docx
    from core import deepdive
    from core.model.ingest import ingest_artifact
    from core.report import docx_bytes
    from core.report.quality import summary
    from core.report.settings import save
    code = 'def calculate():\n    return 1\n'
    aid = store.add_artifact('policy.py', language='Python', transcription=code)
    ingest_artifact(store, None, aid)
    store.set_validation(aid, 'local fixture check', True, '')
    store.set_meta('deepdive', {str(aid): {'name': 'policy.py', 'hash': deepdive._hash(code),
                   'purpose': 'Test calculation', 'facts': [], 'unknowns': [], 'unverifiable': [],
                   'ran_at': '2026-10-01', 'prompt_version': deepdive.PROMPT_VERSION}})
    save(store, {'signed_off': 'yes', 'technical_reviewer': 'Tech', 'business_owner': 'Business', 'it_reviewer': 'IT'})
    from core.report.quality import measure
    original = tmp_path/'policy.py';original.write_text(code)
    measure(store, {'policy.py': original})
    assert not summary(store)['blockers']
    document = docx.Document(io.BytesIO(docx_bytes(store, rescan=False)))
    text = '\n'.join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    assert 'v1.0 Final' not in text
    assert 'Draft for review' in text
    assert 'Blocks issue:' in text
    assert 'open items block issue of the report as final' in text


def test_html_and_editable_diagrams_ignore_redaction_settings(store):
    import io
    import docx
    from core.model.ingest import ingest_artifact
    from core.report import package
    from core.report.settings import save
    aid = store.add_artifact('SECRETCLIENT.py', language='Python', transcription='def calculate():\n    return 1\n')
    ingest_artifact(store, None, aid)
    save(store, {'client': 'Example Agency', 'redact_terms': 'SECRETCLIENT'})
    pkg = package(store, pdf=False)
    output = Path(pkg['dir'])
    word = docx.Document(output / f'{store.info["slug"]}_assessment_report.docx')
    text = '\n'.join([p.text for p in word.paragraphs] + [c.text for t in word.tables for r in t.rows for c in r.cells])
    assert 'SECRETCLIENT' not in text
    assert 'SECRETCLIENT' not in (output / f'{store.info["slug"]}_assessment_report.html').read_text()
    assert 'SECRETCLIENT' not in (output / f'{store.info["slug"]}_diagrams.drawio').read_text()


def test_old_program_response_overwrites_newly_selected_program():
    import subprocess
    source = Path('src/webapp/static/app.js').read_text()
    signature = 'function _programRenderSignature(' + source.split('function _programRenderSignature(', 1)[1].split('\nasync function loadProgram', 1)[0]
    load = 'async function loadProgram(' + source.split('async function loadProgram(', 1)[1].split('\nlet _progPoll', 1)[0]
    script = signature + load + r'''
const assert = require('node:assert/strict');
let _program = 'A', _progSig = '';
const _localRebuilds = new Map(), elements = {};
const $ = id => elements[id] ||= {style: {}, textContent: '', innerHTML: ''};
const document = {querySelectorAll: () => []};
const renderProgress = () => {}, scheduleProgramPoll = () => {}, renderCaptureBar = () => {};
const _fileRow = row => row.name;
const deferred = {};
const _json = url => new Promise(resolve => deferred[url] = resolve);
const data = name => ({program: {}, progress: {total: 1}, artifacts: [{id: 1, name, state: 'done'}], coverage: {missing: []}});
(async () => {
  const oldRequest = loadProgram(true);
  _program = 'B';
  const newRequest = loadProgram(true);
  deferred['/api/programs/B'](data('B file'));
  await newRequest;
  assert.equal($('progFiles').innerHTML, 'B file');
  deferred['/api/programs/A'](data('A file'));
  await oldRequest;
  assert.equal(_program, 'B');
  assert.equal($('progFiles').innerHTML, 'B file');
})().catch(e => { console.error(e); process.exitCode = 1; });
'''
    subprocess.run(['node', '-e', script], check=True, capture_output=True, text=True)


def test_report_list_refresh_stops_when_event_window_reaches_sixty():
    import subprocess
    source = Path('src/webapp/static/app.js').read_text()
    poll = 'async function pollStatus(' + source.split('async function pollStatus(', 1)[1].split('\nfunction setRunning', 1)[0]
    script = poll + r'''
const assert = require('node:assert/strict');
let lastEventCount = 0, _program = '', refreshes = 0, offset = 0;
const events = () => Array.from({length: 60}, (_, i) => ({t: offset + i, kind: 'info', msg: 'event'}));
const fetch = async () => ({json: async () => ({running: false, events: events()})});
const renderStatus = () => {}, renderFlow = () => {}, setRunning = () => {}, stopAnalyzing = () => {};
const loadReports = () => refreshes++;
(async () => {
  await pollStatus();
  assert.equal(refreshes, 1);
  offset = 60;
  await pollStatus();
  assert.equal(refreshes, 2); // new events refresh even when count stays bounded
})().catch(e => { console.error(e); process.exitCode = 1; });
'''
    subprocess.run(['node', '-e', script], check=True, capture_output=True, text=True)


def test_failed_capture_does_not_persist_completed_provider_spend(store, tmp_path, monkeypatch):
    import hotkey_capture as hk
    from core import status, notify
    session = tmp_path / 'session'
    session.mkdir()
    shot = image(session)
    aid = store.add_pending_capture([shot])
    app = hk.App(None)
    app.program = store.info['slug']
    monkeypatch.setattr(status, 'publish', lambda *args, **kwargs: None)
    monkeypatch.setattr(notify, 'notify', lambda *args, **kwargs: None)
    monkeypatch.setattr(app, '_share_copybooks', lambda: None)
    def fail_after_paid_read(*args, **kwargs):
        app.tracker.record('transcribe', 'claude-sonnet-5',
                           SimpleNamespace(input_tokens=1000, output_tokens=100), 1)
        raise RuntimeError('failure after a completed synthetic paid read')
    monkeypatch.setattr(hk, 'run_team_fast', fail_after_paid_read)
    try:
        app._analyse_burst(session, artifact_id=aid, kind='code')
        assert store.artifact(aid)['status'] == 'failed'
        assert app.tracker.total_cost > 0
        assert store.usage()['cost'] == app.tracker.total_cost
        assert len(store.runs()) == 1
    finally:
        app._pool.shutdown(wait=True)


def test_usage_is_durable_before_failure_and_deduplicated(store):
    from core.usage import program_client
    response = SimpleNamespace(model='claude-sonnet-5', usage=SimpleNamespace(input_tokens=1000, output_tokens=100),
                               stop_reason='end_turn')
    client = program_client(store.info['slug'], SimpleNamespace(messages=SimpleNamespace(create=lambda **kw: response)))
    client.messages.create(model='claude-sonnet-5')
    assert len(store.runs()) == 1
    record = client.tracker.take(everything=True)[0]
    store.log_usage_record(record)
    assert len(store.runs()) == 1
    assert store.usage()['cost'] == client.tracker.total_cost
    store.set_meta('budget_usd', store.usage()['cost'])
    from core.usage import BudgetExceeded
    with pytest.raises(BudgetExceeded):
        client.messages.create(model='claude-sonnet-5')
    assert len(store.runs()) == 1


def test_legacy_source_collision_recovered_from_database(store):
    one = store.add_artifact('a.rexx', transcription='say 1')
    two = store.add_artifact('b.rexx', transcription='say 2')
    path = store.artifact(one)['source_path']
    with store.transaction() as db:
        db.execute('UPDATE artifact SET source_path=? WHERE id=?', (path, two))
    with ProgramStore.open(store.info['slug']) as reopened:
        a, b = reopened.artifact(one), reopened.artifact(two)
        assert a['source_path'] != b['source_path']
        assert (reopened.sources_dir / a['source_path']).read_text() == 'say 1'
        assert (reopened.sources_dir / b['source_path']).read_text() == 'say 2'


def test_signoff_invalidated_by_source_change(store):
    from core.report.settings import save, get
    aid = store.add_artifact('policy.py', transcription='print(1)')
    save(store, {'signed_off': 'yes'})
    assert get(store)['signed_off_basis'] == store.model_stamp()
    store.fill_artifact(aid, artifact_type='code', language='Python', transcription='print(2)')
    assert get(store)['signed_off_basis'] != store.model_stamp()


def test_export_zip_redacts_every_text_member_including_nested_office_files(store):
    import io
    import zipfile
    from core.model.ingest import ingest_artifact
    from core.report import package
    from core.report.settings import save
    from core.diagrams.export import bundle
    aid = store.add_artifact('SECRETCLIENT.py', language='Python', transcription='def calculate():\n    return 1\n')
    ingest_artifact(store, None, aid)
    save(store, {'redact_terms': 'SECRETCLIENT'})
    def check(blob):
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            for name in archive.namelist():
                assert 'SECRETCLIENT' not in name
                data = archive.read(name)
                if name.endswith(('.docx', '.vsdx')):
                    check(data)
                elif name.endswith(('.html', '.drawio', '.xml', '.rels', '.json', '.svg')):
                    assert b'SECRETCLIENT' not in data, name
    check(package(store, pdf=False)['zip'])
    check(bundle(store, with_png=False))


def test_usage_callback_can_save_during_capture_ingestion(store, tmp_path, monkeypatch):
    import hotkey_capture as hk
    from core import model, status
    response = SimpleNamespace(model='claude-sonnet-5', usage=SimpleNamespace(input_tokens=1000, output_tokens=100),
                               stop_reason='end_turn')
    app = hk.App(SimpleNamespace(messages=SimpleNamespace(create=lambda **kw: response)))
    app.program = store.info['slug']
    app._attach_usage_persistence()
    aid = store.add_pending_capture([image(tmp_path)])
    store.claim_pending(aid, app._owner)
    def complete(st, client, artifact_id, report):
        client.messages.create(model='claude-sonnet-5')
        st.set_status(artifact_id, 'structured')
        st.clear_pending(artifact_id)
        return artifact_id
    monkeypatch.setattr(model, 'complete_capture', complete)
    monkeypatch.setattr(app, '_analyse_code', lambda *a: None)
    monkeypatch.setattr(app, '_report_recapture', lambda *a: None)
    monkeypatch.setattr(status, 'publish', lambda *a, **kw: None)
    try:
        app._ingest_into_program([], SimpleNamespace(last_report={}), artifact_id=aid)
        assert store.artifact(aid)['status'] == 'structured'
        assert store.usage()['cost'] > 0
        assert len(store.runs()) == 1
    finally:
        app._pool.shutdown(wait=True)
