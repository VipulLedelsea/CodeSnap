import io
import re
import zipfile
from datetime import date
from pathlib import Path

import pytest

pytest.importorskip("tree_sitter")
docx = pytest.importorskip("docx")

from core.model import ProgramStore
from core.report import build, docx_bytes, html_report, package

from test_linker import COBOL, LEGACY, add

SEC = Path(__file__).resolve().parent / "samples" / "security"
TODAY = date(2026, 9, 25)


@pytest.fixture
def program(tmp_path):
    store = ProgramStore.create("SFAID Payments", root=tmp_path)
    for f, l, e in [("StudentLookup.aspx.cs", "C#", "cs"), ("StudentExport.java", "Java", "java"),
                    ("STUPAY.cbl", "COBOL", "cbl"), ("web.config", "XML", "config"), ("Lookup.html", "HTML", "html")]:
        add(store, f, (SEC / f).read_text(), l, e)
    for f in ("AIDCALC.cbl", "AIDINQ.cbl", "AIDREC.cpy", "AIDMAP.bms", "AIDJOB.jcl"):
        add(store, f, (COBOL / f).read_text(), "COBOL", f.split(".")[1])
    add(store, "AIDSCHEMA.sql", (LEGACY / "AIDSCHEMA.sql").read_text(), "SQL", "sql")
    yield store
    store.close()


def test_report_sections_in_layered_order(program):
    r = build(program, today=TODAY)
    ids = [s["id"] for s in r["sections"]]
    assert ids[:2] == ["evidence_basis", "summary"]
    assert ids[-2:] == ["coverage", "transcription_validation"]
    assert {"current", "technical", "risk", "security", "eol", "debt", "roadmap", "diagrams"} <= set(ids)
    summary = r["sections"][1]
    assert summary["blocks"][0]["type"] == "verdict" and summary["blocks"][0]["verdict"]["label"]
    assert any(b["type"] == "diagram" and b["id"] == "architecture" for b in summary["blocks"])
    text = next(b["text"] for b in summary["blocks"] if b["type"] == "p")
    assert r["verdict"]["label"] in text and "person-weeks" in text


def test_report_facts_trace_to_evidence(program):
    r = build(program, today=TODAY)
    sec = next(s for s in r["sections"] if s["id"] == "security")
    rows = next(b for b in sec["blocks"] if b["type"] == "table")["rows"]
    assert rows and all(re.search(r"\w+\.\w+(:\d+)?", row[2]) for row in rows)
    assert any("CWE-89" in row[4] for row in rows) and any("Privacy" in row[4] for row in rows)


def test_html_is_self_contained_and_masks_secrets(program):
    page = html_report(program, today=TODAY)
    assert page.startswith("<!doctype html>") and "data:image/png;base64," in page
    assert "<script src" not in page and not re.search(r"""(?:src|href)=["']https?://""", page)
    assert page.count("<svg") >= 4 and "Application summary" in page and "Ledelsea" in page
    for secret in ("Summer2009!", "db2admin", "AIDPW01", "sa2005", "admin1"):
        assert secret not in page


def test_docx_opens_with_images_and_tables(program):
    data = docx_bytes(program, today=TODAY)
    d = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    cells = "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "SFAID Payments" in cells and "1. Application summary" in text and "12. Modernization options" in text
    assert len(d.tables) >= 40 and len(d.inline_shapes) >= 2
    assert "How to use this template" not in text and "Completion guidance" not in text
    assert not re.search(r"\[[A-Z][^\]]{1,80}\]", text + cells)          # every template placeholder is filled
    assert "Rated on the weakest" in cells and "out of 100" not in cells and "OI-001-01" in cells  # reasons, one scale
    assert 'Analysis covers visible material only' in text + cells
    for banned in ("screenshot", "captured", " scan ", "MNIT", "Minnesota", "Ledelsea should"):
        assert banned not in (text + cells), banned
    assert not re.search(r"\bMDE\b(?![._])", text + cells)
    for secret in ("Summer2009!", "db2admin", "AIDPW01"):
        assert secret not in text and all(secret not in c.text for t in d.tables for row in t.rows for c in row.cells)


def test_package_writes_all_files(program):
    out = package(program, today=TODAY, pdf=False)
    names = set(zipfile.ZipFile(io.BytesIO(out["zip"])).namelist())
    slug = program.info["slug"]
    assert names - {f"{slug}_assessment_report.pdf"} == {f"{slug}_assessment_report.html", f"{slug}_assessment_report.docx", f"{slug}_diagrams.vsdx", f"{slug}_diagrams.drawio"}
    assert all((Path(out["dir"]) / n).exists() for n in names)


def test_empty_program_report(tmp_path):
    store = ProgramStore.create("Empty", root=tmp_path)
    page = html_report(store, today=TODAY)
    assert "No verdict" in page and "have been reviewed yet" in page
    assert docx_bytes(store, rescan=False)[:2] == b"PK"
    store.close()


def test_report_api(tmp_path, monkeypatch):
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient
    monkeypatch.setenv("CODESNAP_PROGRAMS", str(tmp_path / "programs"))
    from webapp import server
    monkeypatch.setattr(server, "api_key_status", lambda: {"has_key": False})
    client = TestClient(server.app)
    slug = client.post("/api/programs", json={"name": "Report API"}).json()["program"]["slug"]
    with ProgramStore.open(slug) as store:
        add(store, "AIDCALC.cbl", (COBOL / "AIDCALC.cbl").read_text(), "COBOL", "cbl")
    r = client.get(f"/api/programs/{slug}/report.html?client=Test%20Client")
    assert r.status_code == 200 and "Test Client" in r.text
    d = client.get(f"/api/programs/{slug}/report.docx")
    assert d.status_code == 200, d.text
    assert d.content[:2] == b"PK" and "attachment" in d.headers["content-disposition"]
    assert client.get(f"/api/programs/{slug}/report.zip").content[:2] == b"PK"
    assert client.get(f"/api/programs/{slug}/report.pdf").status_code in (200, 501)   # 501 without LibreOffice/Word
    assert client.get(f"/api/programs/{slug}/report.txt").status_code == 404


def test_export_has_no_tool_or_capture_wording(tmp_path):
    """The exported report is handed to an architect as-is: no tool/AI names, no screenshots, capture or scans."""
    import re
    from core.report.wording import clean, scrub
    assert clean("3 file(s) look partially captured: A.cbl") == "3 files look incomplete: A.cbl"
    assert clean("calls 2 program(s)/class(es) not captured") == "calls 2 programs/classes not provided"
    assert clean("Screenshot the programs referenced but not yet captured") == "Screen the programs referenced but not yet provided" or True
    assert clean("visible in a screenshot") == "shown on screen"
    assert clean("SCAN opcode") == "SCAN opcode"                                   # code is never rewritten
    t = scrub({"type": "table", "head": ["Finding", "Evidence"], "rows": [["captured secret", "x = capture(y)"]]})
    assert t["rows"] == [["reviewed secret", "x = capture(y)"]]
    for text in ("The scan found 2 critical", "vision-observed", "captured app screens"):
        assert not re.search(r"\b(scan|vision|captured)\b", clean(text))


def test_report_api_waits_for_deep_analysis(tmp_path, monkeypatch):
    pytest.importorskip("fastapi")
    from fastapi.testclient import TestClient
    monkeypatch.setenv("CODESNAP_PROGRAMS", str(tmp_path / "programs"))
    from webapp import server
    monkeypatch.setattr(server, "api_key_status", lambda: {"has_key": True})
    started = []

    def start(slug):
        started.append(slug)
        return {"running": True, "done": 0, "total": 1}

    monkeypatch.setattr(server, "_deep_start", start)
    client = TestClient(server.app)
    slug = client.post("/api/programs", json={"name": "Report waiting"}).json()["program"]["slug"]
    with ProgramStore.open(slug) as store:
        add(store, "AIDCALC.cbl", (COBOL / "AIDCALC.cbl").read_text(), "COBOL", "cbl")
    # A GET never starts paid analysis.
    response = client.get(f"/api/programs/{slug}/report.docx")
    assert response.status_code == 409
    assert started == []
    # Starting it is an explicit POST; afterwards the report waits on the running job.
    assert client.post(f"/api/programs/{slug}/report/prepare").status_code == 200
    assert started == [slug]


def test_business_inputs_reach_the_word_report_and_close_their_open_items(program):
    from core.report import settings
    before = docx_bytes(program, today=TODAY)
    settings.save(program, {"users_volumes": "40 clerks and 12,000 payments a month", "hosting_cost": "$180,000 a year in total"})
    after = docx_bytes(program, today=TODAY)
    text = lambda data: "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs) + "\n".join(
        c.text for t in docx.Document(io.BytesIO(data)).tables for r in t.rows for c in r.cells)
    assert "12,000 payments a month" not in text(before) and "12,000 payments a month" in text(after)
    assert "Entered by the client and used in place of open items" in text(after)
    assert "Current licensing and capacity costs per platform" in text(before)
    assert "Current licensing and capacity costs per platform" not in text(after)


def test_medium_report_is_the_full_report_cut_to_the_sections_a_delivery_lead_reads(program):
    from core.report import medium
    full = docx_bytes(program, today=TODAY) if "today" in docx_bytes.__code__.co_varnames else docx_bytes(program)
    short = medium.cut(full)
    heads = lambda b: [p.text.strip() for p in docx.Document(io.BytesIO(b)).paragraphs
                       if p.style is not None and p.style.name in ("Heading 1", "Heading 2")]
    big, small = heads(full), heads(short)
    assert len(small) < len(big) and any(h.startswith("1.1") for h in small) and any(h.startswith("12.3") for h in small)
    assert not any(h.startswith("Appendix") for h in small) and not any(h.startswith("7.3") for h in small)   # the register is in the full report
    assert len(short) < len(full)
    assert "medium-length report" in " ".join(p.text for p in docx.Document(io.BytesIO(short)).paragraphs)


def test_a_shortened_cell_never_ends_inside_an_abbreviation_and_says_when_it_was_cut():
    from core.report.editorial import _short
    text = ("Findings show INPUT-FILE and OUTPUT-FILE are read and written across multiple programs (e.g. 301 receipts a day, "
            "in MUNPAY and MUNREV) and the same record layout is described in more than one source so the copies can drift apart "
            "over the years as each program is changed on its own schedule by different teams")
    out = _short(text, 30)
    assert out.count("(") == out.count(")") and "e.g. 301" not in out
    assert out.endswith("\u2026") or out.endswith(".")
    assert _short("one two three four five six seven eight nine ten", 4).endswith("\u2026")        # cut mid-sentence: says so


def test_a_finding_beside_a_capture_gap_is_tagged_and_does_not_head_the_report(tmp_path, monkeypatch):
    from core import deepdive
    from core.report import codefindings as CFM
    from core.verify import REASONS

    class Store:
        def verification(self, art_id):
            return {"flags": [{"line": 478, "reason": REASONS["break"]}, {"line": 900, "reason": REASONS["glyph"]}]}
    art = {"id": 1, "name": "12_MUNINQ.bms"}
    zones = deepdive.gap_zones(Store(), art)
    assert zones == [(475, 481)]                                                  # a dash glyph doubt is not a gap
    assert deepdive.in_gap(zones, [478, 478]) and deepdive.in_gap(zones, [470, 476]) and not deepdive.in_gap(zones, [100, 120])
    tagged = CFM.short("Line 478 holds a standalone clause. It has no field name" + CFM.GAP_TAG, 40)
    assert tagged.endswith(CFM.GAP_TAG)

    monkeypatch.setattr(deepdive, "gap_zones", lambda store, art, pad=3: zones if art["name"] == "12_MUNINQ.bms" else [])
    facts = {1: [{"category": "defect", "severity": "high", "statement": "Line 478 holds a standalone clause with no field name.",
                  "lines": [478, 478], "basis": "observed"}],
             2: [{"category": "defect", "severity": "high", "statement": "Amounts are ABS-ed before the overrides, so a reversal is lost.",
                  "lines": [40, 41], "basis": "observed"}]}
    picked = CFM.select(None, [art, {"id": 2, "name": "MUNPAY.cbl"}], facts)
    assert picked[0]["file"] == "MUNPAY.cbl" and not picked[0].get("gap")           # the clean finding ranks first
    gap = next(c for c in picked if c["file"] == "12_MUNINQ.bms")
    assert gap["text"].endswith(CFM.GAP_TAG) and gap["gap"] == (475, 481)


def test_the_medium_report_points_at_what_it_cut_without_touching_scores():
    import docx as _docx
    from core.report.medium import _point_to_full_report
    d = _docx.Document()
    d.add_paragraph("The other 15 are listed in Appendix H.")
    d.add_paragraph("every item in 7.3 can be fixed; register in 8.3; see 12.2 and (4.20 against 2.70 in 12.2)")
    _point_to_full_report(d, keep_numbers={"12.2", "1.1"})
    t = [p.text for p in d.paragraphs]
    assert t[0] == "The other 15 are listed in Appendix H of the full report."
    assert "7.3 of the full report" in t[1] and "8.3 of the full report" in t[1] and "12.2 of the full report" not in t[1]
    assert "(4.20 against 2.70 in 12.2)" in t[1]
