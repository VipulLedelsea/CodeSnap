import io
from pathlib import Path

import pytest

from PIL import Image, ImageDraw
from fastapi.testclient import TestClient

from core.capture import content_changed
from core.model import ProgramStore


def png(image):
    out = io.BytesIO()
    image.save(out, format='PNG')
    return out.getvalue()


def test_scrolled_text_is_kept_but_caret_is_ignored():
    first = Image.new('RGB', (600, 400), 'white')
    second = first.copy()
    for image, offset in [(first, 0), (second, 20)]:
        draw = ImageDraw.Draw(image)
        for line in range(25):
            draw.text((20, line * 15), f'{line + offset:04d}  MOVE PAYMENT-AMOUNT TO OUTPUT-RECORD', fill='black')
    assert content_changed(None, png(first))
    assert not content_changed(png(first), png(first))
    assert content_changed(png(first), png(second))
    caret = first.copy()
    ImageDraw.Draw(caret).line((25, 20, 25, 30), fill='black')
    assert not content_changed(png(first), png(caret))


def test_remove_clears_analysis_and_keeps_other_files(tmp_path, monkeypatch):
    monkeypatch.setenv('CODESNAP_PROGRAMS', str(tmp_path / 'programs'))
    from webapp import server
    with ProgramStore.create('Removal test') as store:
        slug = store.info['slug']
        deleted = store.add_artifact('bad.cbl', transcription='bad capture')
        kept = store.add_artifact('good.cbl', transcription='good capture')
        store.add_finding('security', 'high', 'bad finding', target_type='artifact', target_id=deleted)
        store.add_finding('security', 'low', 'good finding', target_type='artifact', target_id=kept)
        store.set_meta('deepdive', {str(deleted): {'facts': ['bad']}, str(kept): {'facts': ['good']}})
        store.set_meta('assessment', {'stale': True})
        store.set_meta('deepdive_program', {'stale': True})
        store.set_meta('recapture_target', {'artifact_id': deleted, 'name': 'bad.cbl'})
        source = store.sources_dir / 'bad.cbl.v1'
    client = TestClient(server.app)
    url = f'/api/programs/{slug}/artifacts/{deleted}'
    assert client.delete(url).status_code == 200
    assert client.delete(url).status_code == 404
    with ProgramStore.open(slug) as store:
        assert store.artifact(deleted) is None
        assert store.artifact(kept)
        assert [f['title'] for f in store.findings()] == ['good finding']
        assert str(deleted) not in store.get_meta('deepdive')
        assert store.get_meta('assessment') is None
        assert store.get_meta('deepdive_program') is None
        assert store.get_meta('recapture_target') is None
        assert not source.exists()


def test_remove_waits_for_active_review(tmp_path, monkeypatch):
    monkeypatch.setenv('CODESNAP_PROGRAMS', str(tmp_path / 'programs'))
    from webapp import server
    with ProgramStore.create('Busy removal') as store:
        slug = store.info['slug']
        aid = store.add_artifact('busy.cbl')
    monkeypatch.setitem(server._DEEP, slug, {'running': True})
    assert TestClient(server.app).delete(f'/api/programs/{slug}/artifacts/{aid}').status_code == 409
    with ProgramStore.open(slug) as store:
        assert store.artifact(aid)


def test_removal_purges_versions_queue_photos_and_report(tmp_path, monkeypatch):
    from core.model.removal import remove_file
    from core.report import html_report, docx_bytes
    import docx
    monkeypatch.setenv('CODESNAP_PROGRAMS', str(tmp_path / 'programs'))
    image = tmp_path / 'shot.png'
    Image.new('RGB', (40, 20), 'red').save(image)
    with ProgramStore.create('Purge test') as store:
        evidence_id = store.add_evidence(image)
        ev = store.evidence(evidence_id)
        old = store.add_artifact('REMOVED.cbl', transcription='IDENTIFICATION DIVISION.\nPROGRAM-ID. REMOVED.', evidence_ids=[evidence_id])
        current = store.add_artifact('REMOVED.cbl', transcription='IDENTIFICATION DIVISION.\nPROGRAM-ID. REMOVED.', evidence_ids=[evidence_id])
        queued = store.add_pending_capture([image], recapture_of='REMOVED.cbl')
        kept = store.add_artifact('KEPT.cbl', transcription='IDENTIFICATION DIVISION.\nPROGRAM-ID. KEPT.')
        cache = store.text_cache_dir / (ev['sha256'] + '.md')
        cache.write_text('REMOVED capture text')
        exports = store.exports_dir / 'report'
        exports.mkdir()
        (exports / 'old.html').write_text('REMOVED.cbl')
        store.set_meta('deepdive', {str(current): {'facts': [{'statement': 'REMOVED finding'}]}})
        removed = remove_file(store, current)
        assert {old, current, queued} == set(removed)
        assert [a['id'] for a in store.artifacts(False)] == [kept]
        assert store.evidence(evidence_id) is None
        assert not Path(ev['abs_path']).exists()
        assert not cache.exists()
        assert not list(exports.glob('*'))
        assert not store.pending_captures()
        assert not store.claim_pending(queued, '123:worker')
        assert image.exists()  # original upload remains outside managed storage
        assert 'REMOVED' not in html_report(store)
        document = docx.Document(io.BytesIO(docx_bytes(store, rescan=False)))
        text = '\n'.join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
        assert 'REMOVED' not in text


def test_shared_photo_stays_only_with_remaining_file(tmp_path, monkeypatch):
    from core.model.removal import remove_file
    monkeypatch.setenv('CODESNAP_PROGRAMS', str(tmp_path / 'programs'))
    with ProgramStore.create('Shared photo') as store:
        evidence_id = store.add_evidence(png(Image.new('RGB', (40, 20), 'green')))
        deleted = store.add_artifact('bad.cbl', evidence_ids=[evidence_id])
        kept = store.add_artifact('good.cbl', evidence_ids=[evidence_id])
        remove_file(store, deleted)
        assert store.artifact_evidence(deleted) == []
        assert store.artifact_evidence(kept)[0]['id'] == evidence_id
        assert Path(store.evidence(evidence_id)['abs_path']).exists()


def test_queued_worker_cannot_claim_removed_file(tmp_path, monkeypatch):
    from core.model.removal import remove_file
    monkeypatch.setenv('CODESNAP_PROGRAMS', str(tmp_path / 'programs'))
    image = tmp_path / 'shot.png'
    Image.new('RGB', (40, 20), 'blue').save(image)
    with ProgramStore.create('Queue removal') as store:
        aid = store.add_pending_capture([image])
        slug = store.info['slug']
        with ProgramStore.open(slug) as worker:
            assert aid in worker.pending_captures()
            remove_file(store, aid)
            assert not worker.claim_pending(aid, '123:worker')


@pytest.fixture(autouse=True)
def isolated_scratch(tmp_path, monkeypatch):
    monkeypatch.setattr("core.model.removal.CAPTURE_ROOT", tmp_path / "captures")


def test_removal_deletes_scratch_and_only_its_text_cache(tmp_path, monkeypatch):
    from core.model.removal import remove_file, CAPTURE_ROOT
    session = CAPTURE_ROOT / 'session_test'
    session.mkdir(parents=True)
    shot = session / '001.png'
    Image.new('RGB', (40, 20), 'purple').save(shot)
    with ProgramStore.create('Scratch removal', root=tmp_path / 'programs') as store:
        aid = store.add_pending_capture([shot])
        digest = store.artifact_evidence(aid)[0]['sha256']
        cache = session / '.cache'
        cache.mkdir()
        owned = cache / (digest + '.md')
        owned.write_text('removed text')
        kept = cache / 'other.md'
        kept.write_text('remaining text')
        remove_file(store, aid)
        assert not shot.exists()
        assert not owned.exists()
        assert kept.exists()


def test_report_changed_during_build_cannot_be_cached(tmp_path, monkeypatch):
    from core import report as R
    def changed_package(store, **kwargs):
        output = store.exports_dir / 'report'
        output.mkdir(parents=True, exist_ok=True)
        (output / 'stale.html').write_text('removed file')
        store.add_artifact('new.cbl')
        return {'zip': b'zip', 'files': ['stale.html'], 'verdict': {}}
    monkeypatch.setattr(R, 'package', changed_package)
    with ProgramStore.create('Concurrent report', root=tmp_path / 'programs') as store:
        with pytest.raises(RuntimeError, match='analysis changed'):
            R.cached_package(store)
        assert not (store.exports_dir / 'report/stale.html').exists()
        assert store.get_meta('report_stamp') is None
