"""The medium-length report: the full assessment report cut down to the sections a delivery lead reads.

The full report (core.report.template_docx) carries every register and appendix; the executive summary (core.report.summary_doc)
is two pages. This keeps the application summary, technical profile, architecture, health, debt, security, risk, modernization options
and open items, trims long tables to their first rows, and drops the appendices and the supporting inventories. It is cut from the same
document, so a figure in it is the same figure as in the full report.
"""
import io
import re
from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn

KEEP_H1 = {"1", "3", "4", "6", "7", "8", "9", "12", "13"}
KEEP_H2 = {"1.1", "1.2", "1.3", "3.1", "3.2", "4.1", "4.4", "6.1", "6.2", "7.1", "7.5", "8.1", "8.2", "9.1",
           "12.1", "12.2", "12.3", "12.4", "13.4", "13.5"}
KEEP_NAMED_H1 = ("document control",)
MAX_ROWS = 12                                  # rows of a long table kept (the header row included)
NOTE = ("This is the medium-length report. It carries the summary, profile, architecture, health, debt, security, risk, options and "
        "open items. The full report adds every register, inventory and appendix, and the executive summary is the two-page version.")


def _style(p):
    name = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return name.get(qn("w:val")) if name is not None else ""


def _text(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _number(text):
    m = re.match(r"^(\d+(?:\.\d+)?)\.?\s", text.strip() + " ")
    return m.group(1) if m else ""


_APPENDIX = re.compile(r"\bAppendix ([A-I])\b(?! of the full report)")
_SECTION_REF = re.compile(r"(?<![\d.\w-])((?:Sections? |in |see |\()(\d{1,2}\.[1-9]))(?![\d]|\.\d|\w)")


def _point_to_full_report(doc, keep_numbers):
    """What the medium report cut away is still in the full report: a reference to it says so instead of pointing at nothing."""
    def fix(text):
        text = _APPENDIX.sub(r"Appendix \1 of the full report", text)
        def section(m):
            return m.group(0) if m.group(2) in keep_numbers else m.group(0) + (" of the full report" if not m.group(1).startswith("(") else " (full report)")
        return _SECTION_REF.sub(section, text)
    paragraphs = list(doc.paragraphs) + [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
    for para in paragraphs:
        for run in para.runs:
            new = fix(run.text)
            if new != run.text:
                run.text = new


def cut(full: bytes) -> bytes:
    doc = Document(io.BytesIO(full))
    body = doc.element.body
    keep_h1, keep_h2, dropped_titles = False, True, set()
    seen_heading = False
    note_at = None
    plain = next((deepcopy(p._p) for p in doc.paragraphs if _style(p._p) in ("", "Normal") and len(p.text) > 60), None)
    for el in list(body):
        tag = el.tag
        if tag == qn("w:sectPr"):
            continue
        style = _style(el) if tag == qn("w:p") else ""
        if style == "Heading1":
            seen_heading = True
            title = _text(el).strip()
            n = _number(title)
            keep_h1 = n in KEEP_H1 or title.lower().startswith(KEEP_NAMED_H1)
            keep_h2 = True
            if keep_h1 and n == "1" and plain is not None and note_at is None:
                note_at = el
            if not keep_h1:
                dropped_titles.add(title)
                body.remove(el)
            continue
        if style == "Heading2":
            seen_heading = True
            title = _text(el).strip()
            keep_h2 = keep_h1 and _number(title) in KEEP_H2
            if not keep_h2:
                body.remove(el)
            continue
        if not seen_heading:
            continue                                   # the cover, document title and contents lines are handled below
        if not (keep_h1 and keep_h2):
            body.remove(el)
            continue
        if tag == qn("w:tbl"):
            rows = el.findall(qn("w:tr"))
            if len(rows) > MAX_ROWS + 1:
                for row in rows[MAX_ROWS:]:
                    el.remove(row)
                if plain is not None:
                    more = deepcopy(plain)
                    for t in list(more.iter(qn("w:t")))[1:]:
                        t.text = ""
                    first = next(more.iter(qn("w:t")))
                    first.text = f"{len(rows) - MAX_ROWS} more rows are in the full report."
                    el.addnext(more)
    _point_to_full_report(doc, keep_numbers=KEEP_H2 | KEEP_H1)
    for el in list(body):                              # contents lines that name a section that is gone
        if el.tag == qn("w:p") and _text(el).strip() in dropped_titles and _style(el) not in ("Heading1", "Heading2"):
            body.remove(el)
    if note_at is not None and plain is not None:
        note = deepcopy(plain)
        texts = list(note.iter(qn("w:t")))
        texts[0].text = NOTE
        for t in texts[1:]:
            t.text = ""
        note_at.addprevious(note)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
