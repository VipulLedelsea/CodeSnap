"""The executive summary as its own Word document, separate from the assessment report.

It is built from the saved summary (core.exec_summary) and only while that summary still matches the findings it was written
from, so the file never describes a program the findings no longer support. Finding tags become file and line references, the
report's wording cleanup and the client redaction rules are applied, and the document says when the review behind it is a draft.
"""
import io
import re
from datetime import datetime

DRAFT_NOTE = "DRAFT: first-pass review, not yet checked line by line."


def _confidence(level, words, zones):
    """The confidence line, corrected when files have gaps: "no reading concerns" cannot stand beside a copy with missing lines."""
    if not zones:
        return level, words
    words = re.sub(r"\s*,?\s*with no (?:reading|capture) concerns,?", "", words, flags=re.I)
    files = ", ".join(sorted(zones)[:4]) + (" and more" if len(zones) > 4 else "")
    note = f"The copy of {len(zones)} file{'s' if len(zones) != 1 else ''} has gaps or cut-off lines ({files}), so findings near them are unconfirmed. "
    return ("low" if level == "high" and len(zones) > 2 else "medium" if level == "high" else level), note + words


def _gap_items(store, redact):
    """One line per file whose captured copy has a gap or cut-off: what the summary says about those lines is unconfirmed."""
    from core import deepdive
    out = []
    for art in store.artifacts():
        try:
            zones = deepdive.gap_zones(store, art)
        except Exception:  # noqa: BLE001
            continue
        if zones:
            spans = ", ".join(f"{a}\u2013{b}" for a, b in zones[:4]) + (" and more" if len(zones) > 4 else "")
            out.append(redact(f"Capture {art['name']} again around lines {spans}: the copy has a gap or cut-off there, so findings near those lines are unconfirmed"))
    return out


def content(store):
    """The summary as plain, finished text ready to lay out, or None when there is no current summary."""
    from core import deepdive
    from core import exec_summary as ES
    summary = ES.current(store)
    if not summary:
        return None
    from . import _redactor
    from .wording import clean
    index = ES.finding_index(store)
    verbatim = [a["name"] for a in store.artifacts()]
    redact = _redactor(store)

    one_line = lambda v: re.sub(r"\blines (\d+)-\1\b", r"line \1", v)     # "lines 2-2" reads better as "line 2"

    def text(value):
        value = ES.plain_refs(value or "", index, limit=3)
        return one_line(redact(clean(ES.clean_style(value), verbatim=verbatim)).strip())

    # a statement that rests on lines beside a gap or cut-off in the copy may describe the capture, not the code: say so
    zones = {}
    for art in store.artifacts():
        try:
            z = deepdive.gap_zones(store, art)
        except Exception:  # noqa: BLE001
            z = []
        if z:
            zones[art["name"]] = z
    cite = re.compile(r"([\w.\-]+\.\w{2,6}) lines? (\d+)(?:\s*[-\u2013]\s*(\d+))?")

    def beside_gap(*texts):
        for t in texts:
            for m in cite.finditer(t or ""):
                a = int(m.group(2)); b = int(m.group(3) or a)
                if any(a <= hi and b >= lo for lo, hi in zones.get(m.group(1), [])):
                    return True
        return False
    UNCONFIRMED = " (unconfirmed: the copy has a gap or cut-off beside these lines)"
    mark = lambda t, *also: t + UNCONFIRMED if beside_gap(t, *also) else t

    order = {k: i for i, k in enumerate(ES.SEVERITIES)}
    risks = sorted(summary.get("risks") or [], key=lambda r: order.get(r.get("severity"), 9))
    where = lambda r: one_line(redact(ES._where(r.get("facts") or [], index, limit=3).strip(" ()")))
    ran = summary.get("ran_at") or ""
    try:
        date = datetime.strptime(ran[:10], "%Y-%m-%d").strftime("%d %B %Y").lstrip("0")
    except ValueError:
        date = datetime.now().strftime("%d %B %Y").lstrip("0")
    slug = store.info["slug"]
    return {
        "title": "Executive summary",
        "program": redact(store.info.get("name") or slug),
        "date": date,
        "draft": DRAFT_NOTE if deepdive.report_stage(store) == "draft" else "",
        "verdict": summary["verdict"],
        "provisional": bool(summary.get("provisional")),
        "needed": [text(x).rstrip(" .") for x in (summary.get("needed_to_finalize") or []) if str(x).strip()] + _gap_items(store, redact),
        "what": mark(text(summary.get("what_it_does"))),
        "reasons": [mark(text(r)) for r in summary.get("reasons") or [] if str(r).strip()],
        "risks": [(r.get("severity") or "Medium", mark(text(r.get("text")), where(r)), where(r)) for r in risks],
        "lifecycle": (bool(summary.get("lifecycle_verified")), text(summary.get("lifecycle"))),
        "decision": text(summary.get("decision")),
        "confidence": _confidence(summary.get("confidence") or "low", text(summary.get("confidence_and_coverage")), zones),
        "evidence": [[redact(c) if isinstance(c, str) else c for c in row] for row in ES.evidence_rows(summary, index)],
        "report_file": f"{slug}_assessment_report.docx",
    }


def docx_bytes(store):
    """The Word file, or None when there is no current summary."""
    c = content(store)
    if c is None:
        return None
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for level, size in (("Title", 24), ("Heading 1", 14)):
        style = doc.styles[level]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

    doc.add_paragraph(c["title"], style="Title")
    sub = doc.add_paragraph(f"{c['program']}  |  {c['date']}")
    sub.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    if c["draft"]:
        p = doc.add_paragraph()
        run = p.add_run(c["draft"])
        run.bold = True
        run.font.color.rgb = RGBColor(0xB0, 0x3A, 0x2E)

    def heading(label):
        doc.add_paragraph(label, style="Heading 1")

    def shade(cell, fill):
        pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        pr.append(shd)

    def table(head, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, label in enumerate(head):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(label)
            run.bold = True
            shade(cell, "E7EBF3")
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)
        if widths:
            t.autofit = False                      # Word and LibreOffice both need fixed widths set on the columns and the cells
            for i, w in enumerate(widths):
                t.columns[i].width = w
                for cell in t.columns[i].cells:
                    cell.width = w
        doc.add_paragraph("")

    heading("Verdict")
    p = doc.add_paragraph()
    p.add_run(c["verdict"] + (" (provisional)" if c["provisional"] else "")).bold = True
    if c["provisional"] and c["needed"]:
        doc.add_paragraph("Needed to finalize this verdict:")
        for item in c["needed"]:
            doc.add_paragraph(item, style="List Bullet")
    if c["what"]:
        heading("What the application does")
        doc.add_paragraph(c["what"])
    if c["reasons"]:
        heading("Why")
        for reason in c["reasons"]:
            doc.add_paragraph(reason, style="List Number")
    if c["risks"]:
        heading("Main risks")
        from docx.shared import Inches
        table(["Severity", "Risk", "Where"], [[s, t, w] for s, t, w in c["risks"]], [Inches(0.9), Inches(3.9), Inches(1.7)])
    heading("Technology lifecycle")
    verified, life = c["lifecycle"]
    doc.add_paragraph(("Verified. " if verified else "Not verified. ") + life)
    heading("Decision needed")
    doc.add_paragraph(c["decision"])
    heading("Confidence and coverage")
    level, why = c["confidence"]
    doc.add_paragraph(f"Confidence {level}. {why}")
    if c["evidence"]:
        heading("Evidence for the findings cited")
        doc.add_paragraph("Each statement above rests on the review findings below.")
        table(["File", "Lines", "Type", "Finding"], c["evidence"])
    note = doc.add_paragraph(f"The full assessment, with the supporting detail, is in {c['report_file']}.")
    note.runs[0].italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

