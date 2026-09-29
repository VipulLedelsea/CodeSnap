import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from core.diagrams import png

from .html import LOGO

INK, MUTED, COPPER = RGBColor(0x2A, 0x25, 0x21), RGBColor(0x6F, 0x65, 0x5C), RGBColor(0xBE, 0x6E, 0x4A)
FILL = {"critical": "FDE2E1", "high": "FDE7DA", "medium": "FFF1D6", "low": "E8F1FB", "info": "EFEFF2",
        "End of life": "FDE2E1", "Legacy — no upgrade path": "FFF1D6", "Extended support only": "FFF1D6",
        "Ends within 12 months": "FFF1D6", "Supported": "E3F5E8", "Version not confirmed": "EFEFF2"}
MATRIX = {"low": "E3F5E8", "medium": "FFF1D6", "high": "FDE2E1", "critical": "F4B4B0"}


def _shade(cell, hexcolor):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tc.append(shd)


def _runs(par, text, size=None, color=None, bold=False):
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", str(text))):
        if not part:
            continue
        r = par.add_run(part)
        r.bold = bold or i % 2 == 1
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return par


def _table(doc, b):
    rows = b["rows"] or [["None."] + [""] * (len(b["head"]) - 1)]
    t = doc.add_table(rows=1, cols=len(b["head"]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    size = 7.5 if b.get("small") else 8.5
    for i, h in enumerate(b["head"]):
        c = t.rows[0].cells[i]
        c.text = ""
        _runs(c.paragraphs[0], h, size, MUTED, True)
        _shade(c, "FBF8F5")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            _runs(cells[i].paragraphs[0], v, size, INK)
            key = str(v).split(" ")[0].lower() if b.get("sev_col") == i else (str(v) if b.get("status_col") == i else None)
            if key and key in FILL:
                _shade(cells[i], FILL[key])
    doc.add_paragraph()


def _diagram(doc, d, width):
    from .content import DIAGRAM_NOTES
    scale = 2.0 if d["width"] < 2400 else 1.4
    max_h = 6.3
    w_in = min(width, d["width"] / 96, max_h * d["width"] / max(d["height"], 1))
    try:
        doc.add_picture(io.BytesIO(png(d, scale=scale)), width=Inches(w_in))
    except Exception as exc:
        _runs(doc.add_paragraph(), f"*Diagram could not be rendered ({type(exc).__name__}); see the Visio file.*", 9, MUTED)
    cap = doc.add_paragraph()
    _runs(cap, f"**{d['title']}.** {DIAGRAM_NOTES.get(d['kind'], '')}", 8, MUTED)


def render(report: dict, diagrams: dict) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.7))
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    width = 11 - 1.4
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10)
    for lvl, size in ((1, 16), (2, 12)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.color.rgb = INK if lvl == 2 else COPPER
        st.font.size = Pt(size)
    cp = doc.core_properties
    cp.author = cp.last_modified_by = report.get("prepared_by") or "Ledelsea"
    cp.title = f"{report['program']} — Application Assessment Report"
    cp.subject = f"Prepared for {report['client']}"
    cp.comments = cp.keywords = cp.category = ""
    footer = sec.footer.paragraphs[0]
    _runs(footer, f"Ledelsea · {report['program']} — Application Assessment Report · confidential — prepared for "
                  f"{report['client']}", 8, MUTED)
    if LOGO.exists():
        doc.add_picture(str(LOGO), width=Inches(2.2))
    title = doc.add_paragraph()
    _runs(title, report["program"], 28, INK, True)
    sub = doc.add_paragraph()
    _runs(sub, "Application Assessment Report", 16, COPPER, True)
    meta = doc.add_paragraph()
    _runs(meta, f"Prepared for {report['client']}\nby {report['prepared_by']} · {report['date']}", 11, MUTED)
    v = report.get("verdict") or {}
    box = doc.add_table(rows=1, cols=4)
    box.style = "Table Grid"
    vals = [("Verdict", f"{v.get('label', 'n/a')} ({v.get('bucket', '')})"),
            ("Overall score", f"{(report['assessment']['scores'].get('overall') or {}).get('score', 'n/a')}/100"),
            ("Total risk", report["assessment"]["total_risk"]["level"].title()),
            ("Confidence", report["assessment"]["confidence"]["level"].title())]
    for c, (k, val) in zip(box.rows[0].cells, vals):
        c.text = ""
        _runs(c.paragraphs[0], k.upper(), 8, MUTED, True)
        _runs(c.add_paragraph(), val, 13, INK, True)
        _shade(c, "F7ECE5")
    doc.add_paragraph()
    toc = doc.add_paragraph()
    _runs(toc, "Contents\n", 11, COPPER, True)
    _runs(toc, "\n".join(f"{i}. {s['title']}" for i, s in enumerate(report["sections"], 1)), 10, INK)
    for i, s in enumerate(report["sections"], 1):
        h = doc.add_heading(f"{i}. {s['title']}", level=1)
        h.paragraph_format.page_break_before = True
        for b in s["blocks"]:
            t = b["type"]
            if t == "p":
                _runs(doc.add_paragraph(), b["text"])
            elif t == "h":
                doc.add_heading(b["text"], level=2)
            elif t == "bullets":
                for item in b["items"]:
                    _runs(doc.add_paragraph(style="List Bullet"), item)
            elif t == "kv":
                tbl = doc.add_table(rows=2, cols=len(b["items"]))
                tbl.style = "Table Grid"
                for ci, (k, val) in enumerate(b["items"]):
                    tbl.rows[0].cells[ci].text = ""
                    _runs(tbl.rows[0].cells[ci].paragraphs[0], k.upper(), 7.5, MUTED, True)
                    tbl.rows[1].cells[ci].text = ""
                    _runs(tbl.rows[1].cells[ci].paragraphs[0], val, 11, INK, True)
                doc.add_paragraph()
            elif t == "table":
                _table(doc, b)
            elif t == "verdict":
                vv = b["verdict"] or {}
                p = doc.add_paragraph()
                _runs(p, f"{vv.get('label', 'No verdict')}", 20, COPPER, True)
                _runs(p, f"   {vv.get('bucket', '').upper()} · confidence {b['confidence']['level']}", 10, MUTED, True)
            elif t == "matrix":
                tbl = doc.add_table(rows=7, cols=6)
                tbl.style = "Table Grid"
                tbl.rows[0].cells[0].text = "Likelihood ↓ / Impact →"
                for ci in range(1, 6):
                    tbl.rows[6].cells[ci].text = str(ci)
                for ri, row in enumerate(b["cells"]):
                    like = 5 - ri
                    tbl.rows[ri + 1].cells[0].text = str(like)
                    for ci, names in enumerate(row):
                        cell = tbl.rows[ri + 1].cells[ci + 1]
                        cell.text = ""
                        if names:
                            _runs(cell.paragraphs[0], f"{len(names)}: " + ", ".join(names[:3]) + ("…" if len(names) > 3 else ""), 7)
                        sc = like * (ci + 1)
                        _shade(cell, MATRIX["critical" if sc >= 20 else "high" if sc >= 12 else "medium" if sc >= 6 else "low"])
                doc.add_paragraph()
            elif t == "diagram" and diagrams.get(b["id"]):
                _diagram(doc, diagrams[b["id"]], width)
            elif t == "diagrams":
                for d in [d for d in diagrams.values() if d["kind"] == b["kind"]][: b.get("limit", 4)]:
                    _diagram(doc, d, width)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
